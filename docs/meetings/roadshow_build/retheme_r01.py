# -*- coding: utf-8 -*-
"""แปลงเด็ค .pptx ที่ทำด้วยธีม TSG R00 (ปก/ปิดพื้นเขียว) → **R01** โดยไม่แตะเนื้อหา (2026-10-12)
ใช้กับเด็คเก่าที่ไม่มีสคริปต์สร้าง: python3 retheme_r01.py <file.pptx> [...]   (เขียนทับไฟล์เดิม)

ทำอะไร:
  · ลบพื้นหลัง R00 (รูป bg_title.png · สี่เหลี่ยมเขียวเข้มเต็มสไลด์) ⇒ พื้นขาว
  · divider R00 (div_*.png) ⇒ divider R01: รูปโรงงานขวา + ระนาบขาวขอบเฉียง + หัวข้อเขียวซ้าย
  · โลโก้ R00 ⇒ โลโก้ R01 (ปก: บนกลาง · footer: ตำแหน่ง R01) · ลายน้ำโลโก้ใหญ่ในหน้าเนื้อหา = ลบ (R01 ไม่มี)
  · สีตามคู่ R00 → R01 — แยก "สีตัวอักษร" กับ "สีพื้น/เส้น" (เขียวเข้ม 0D3D14 ยังเป็นพื้น headline box/หัวตาราง ได้)
  · บนสไลด์ที่เคยพื้นเขียว: ตัวขาว/ทอง ⇒ เขียว/ส้ม (ไม่งั้นมองไม่เห็นบนพื้นขาว)
  · footer + เลขหน้า ย้ายไปตำแหน่ง R01 · หน้าปิดที่ไม่มี footer ใส่ให้
ระบุรูป R00 ด้วย md5 (ห้ามเดาจากขนาด)
"""
import hashlib, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn
import tsg_theme as T

A = T.ASSET
# md5 ของรูป R00 (ไฟล์ถูกลบออกจาก repo แล้ว 12/10 — เก็บไว้แค่ลายนิ้วมือเพื่อจับในเด็คเก่า)
MD5 = {
    '2578c25ce3768bb2d436dfa20006e6d7': 'bg_title.png',
    '02df999eae910ad396e51f1d15a10889': 'div_pd3.png',
    '3f6f9e8f36c03ae1d869d2572d7af0d7': 'div_pd4.png',
    'c1e8e00a4bb6a7bb7b7da432aacd50b1': 'div_jig.png',
    '74963538f1cb3d236a11d2e6171f37c2': 'logo_big.png',
    'caa820954bc0b935e82ddb31a749e091': 'logo_footer.png',
}
R00_BG, R00_LOGO_BIG, R00_LOGO_FOOT = 'bg_title.png', 'logo_big.png', 'logo_footer.png'
R00_DIV = {'div_pd3.png', 'div_pd4.png', 'div_jig.png'}
DIV_PHOTO = ['photo_press2500t.png', 'photo_robots.png', 'photo_hotpress.png']

TEXT_MAP = {'0D3D14': '068734', '2C5F2D': '068734', '14571F': '068734', 'C0561E': 'D95323', 'FD8342': 'D95323'}
TEXT_MAP_DARK = {**TEXT_MAP, 'FFFFFF': '068734', 'EBD9B0': 'D95323', 'CFE0C8': '555555', 'ECF1E9': '068734'}
FILL_MAP = {'C0561E': 'D95323', 'FD8342': 'E2772E', 'ECF1E9': 'F2F2F2', 'D8E4D0': 'D9D9D9', '2C5F2D': '068734', '14571F': '068734'}
TEXT_TAGS = {qn('a:rPr'), qn('a:defRPr'), qn('a:endParaRPr')}


def img_name(sh):
    try:
        return MD5.get(hashlib.md5(sh.image.blob).hexdigest())
    except Exception:
        return None


def is_text_color(el):
    a = el.getparent()
    while a is not None:
        if a.tag in TEXT_TAGS:
            return True
        if a.tag in (qn('p:spPr'), qn('a:tcPr'), qn('a:ln')):
            return False
        a = a.getparent()
    return False


def send_to_back(slide, shape, pos=2):
    tree = slide.shapes._spTree
    tree.remove(shape._element)
    tree.insert(pos, shape._element)


def full(sh, p):
    return sh.width >= p.slide_width * 0.9 and sh.height >= p.slide_height * 0.9


def convert(path):
    p = Presentation(path)
    n = len(p.slides)
    div_i = 0
    for idx, s in enumerate(p.slides, 1):
        dark = False
        divider = False
        for sh in list(s.shapes):
            nm = img_name(sh) if sh.shape_type == 13 else None
            if nm == R00_BG:
                sh._element.getparent().remove(sh._element); dark = True
            elif nm in R00_DIV:
                sh._element.getparent().remove(sh._element); dark = divider = True
            elif sh.shape_type != 13 and full(sh, p):
                try:
                    if sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == '0D3D14':
                        sh._element.getparent().remove(sh._element); dark = True
                except Exception:
                    pass
        for sh in list(s.shapes):
            nm = img_name(sh) if sh.shape_type == 13 else None
            if nm == R00_LOGO_BIG:
                sh._element.getparent().remove(sh._element)
                if dark and not divider:   # ปก: โลโก้ R01 บนกลาง · หน้าเนื้อหา: ลายน้ำ R00 = ลบทิ้ง
                    T.logo(s, 5.92, 0.28, 1.25)
            elif nm == R00_LOGO_FOOT:
                sh._element.getparent().remove(sh._element)
                T.logo(s, 0.273, 7.052, 0.26)
        if divider:
            photo = DIV_PHOTO[div_i % len(DIV_PHOTO)]; div_i += 1
            pic = T._cover_pic(s, os.path.join(A, photo), 5.92, 0, 7.41, 7.5)
            wr = T.rect(s, -0.05, 0, 6.1, 7.5, fill=T.WHITE)
            tri = T.rect(s, 6.05, 0, 1.31, 7.5, fill=T.WHITE, shape=__import__('pptx.enum.shapes', fromlist=['MSO_SHAPE']).MSO_SHAPE.RIGHT_TRIANGLE)
            tri._element.spPr.find(qn('a:xfrm')).set('flipV', '1')
            for k, el in enumerate((pic, wr, tri)):
                send_to_back(s, el, 2 + k)
            for sh in s.shapes:   # หัวข้อ divider R00 อยู่ฝั่งขวา ⇒ R01 ย้ายไปฝั่งซ้าย
                if sh.has_text_frame and sh.text_frame.text.strip() and sh.left > Inches(6):
                    sh.left, sh.width = Inches(0.3), Inches(5.9)
                    for para in sh.text_frame.paragraphs:
                        para.alignment = 2
                        for r in para.runs:
                            r.font.bold = True
        # สี
        cmap = TEXT_MAP_DARK if dark else TEXT_MAP
        for el in s._element.iter(qn('a:srgbClr')):
            v = el.get('val', '').upper()
            m = cmap if is_text_color(el) else FILL_MAP
            if v in m:
                el.set('val', m[v])
        # footer/เลขหน้า → ตำแหน่ง R01
        has_footer = False
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if t == 'THAI SUMMIT GROUP':
                has_footer = True
                sh.left, sh.top, sh.width, sh.height = Inches(0.62), Inches(6.948), Inches(3.55), Inches(0.438)
                for r in sh.text_frame.paragraphs[0].runs:
                    r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = T.GREEN
            elif t.isdigit() and sh.top > Inches(6.8) and sh.left > Inches(10):
                sh.left, sh.top, sh.width, sh.height = Inches(10.28), Inches(7.12), Inches(2.75), Inches(0.32)
                for r in sh.text_frame.paragraphs[0].runs:
                    r.font.color.rgb = T.GREEN
        if idx == n and dark and not has_footer:   # หน้าปิด R01 มี footer
            T.footer(s, idx)
    p.save(path)
    print('R01', os.path.basename(path), 'slides', n)


if __name__ == '__main__':
    for f in sys.argv[1:]:
        convert(f)

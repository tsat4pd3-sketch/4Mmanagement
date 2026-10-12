# -*- coding: utf-8 -*-
"""Thai Summit Group presentation theme — **Revision 01** (python-pptx)
=====================================================================
แหล่งความจริง: "Presentation_template_VX_R01.pptx" (user ส่งมา 2026-08-24) — สเปคถอดไว้แล้วที่
หัวไฟล์ `src/lib/monthlyReviewPptx.js` (ตัวสร้างรายงานเดือนในแอปใช้ R01 อยู่) · ไฟล์นี้ต้องตรงกับนั้นเสมอ

R01 ต่างจาก R00 (ที่ skill `tsg-presentation-theme` ยังเป็น):
  · **พื้นขาวทุกสไลด์ — รวมหน้าปกและหน้าปิด** (เลิกพื้นเขียวเข้ม/gradient ของ R00)
  · เขียวหลัก (หัวเรื่อง/ข้อความ/footer/เลขหน้า) = 068734 · ส้ม accent = D95323
  · เขียวเข้ม 0D3D14 เหลือแค่ Headline box (พื้นเข้มตัวขาว) + หัวตาราง
  · หัวเรื่อง (0.28,0.28) 36 Bold · subtitle (0.42,1.05) 20 Bold · ย่อให้จบบรรทัดเดียวเสมอ
  · footer: โลโก้ R01 (0.273,7.052 · 0.26) + "THAI SUMMIT GROUP" 20 Bold + เลขหน้า (10.28,7.12) 12 ขวา
  · ปก: โลโก้บนกลาง · หัวเรื่อง 40 Bold เขียวกลางสไลด์ · My Quality Declaration (EN+TH) · แถบรูป 4 ใบท้ายสไลด์
  · Agenda: ลิสต์เลขสีเขียว (ไม่มีวงกลมส้มแล้ว)
  · divider: รูปโรงงานฝั่งขวา + ระนาบขาวขอบเฉียง + หัวข้อเขียวฝั่งซ้าย (ไม่มี footer)
  · ปิด: พื้นขาว คำขวัญเขียว 36 Bold + footer
ชื่อค่าคงที่ (GREEN_D/GREEN_M/…) คงไว้ให้สคริปต์เดิมใช้ต่อได้ — ค่าข้างในเป็นของ R01 แล้ว
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import ImageFont

ASSET = os.path.join(os.path.dirname(os.path.abspath(__file__)), "tsg_assets")
REVISION = "R01"

# ---- Palette R01 ----
GREEN = RGBColor(0x06, 0x87, 0x34)      # เขียวหลัก R01 — หัวเรื่อง/ข้อความ/footer/เลขหน้า
GREEN_D = RGBColor(0x0D, 0x3D, 0x14)    # Headline box + หัวตาราง เท่านั้น
GREEN_M = GREEN                         # (ชื่อเดิม R00) ข้อความเนื้อหา = เขียวหลัก
GREEN_TINT = RGBColor(0xF2, 0xF2, 0xF2) # แถวสลับตาราง / กล่องพื้นอ่อน (palette ทางการ R01)
BORDER = RGBColor(0xD9, 0xD9, 0xD9)
ORANGE = RGBColor(0xD9, 0x53, 0x23)     # ส้ม accent R01 — ตัวเลข stat / badge / สถานะ
ORANGE_LT = RGBColor(0xF6, 0xCC, 0xBE)
AMBER = RGBColor(0xC8, 0x8A, 0x00)      # สีที่ 3 ("Other") — สถานะ watch
GREY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = AMBER                            # (ชื่อเดิม R00 — ปก R01 ไม่มีตัวทองแล้ว)
FONT = "Tahoma"
SW, SH = 13.333, 7.5
LOGO = "logo_r01.png"

# ---- วัดความกว้างข้อความ (ใช้ย่อหัวเรื่องให้จบบรรทัดเดียว — ห้ามพึ่ง autofit ของ PowerPoint) ----
_F = {}
_LAT = "/usr/share/fonts/truetype/dejavu/DejaVuSans{}.ttf"
_THA = "/usr/share/fonts/opentype/tlwg/Loma{}.otf"
def _font(kind, bold):
    k = (kind, bold)
    if k not in _F:
        path = (_THA if kind == "th" else _LAT).format("-Bold" if bold else "")
        if not os.path.exists(path):
            path = (_THA if kind == "th" else _LAT).format("")
        _F[k] = ImageFont.truetype(path, 100)
    return _F[k]
def width_in(t, pt, bold=False):
    """ประมาณความกว้าง (นิ้ว) · DejaVu กว้างกว่า Tahoma ~7% ⇒ คูณ 0.93 · ไทยวัดด้วย Loma ×1.05 (Tahoma ไทยกว้างกว่า)"""
    tot = 0.0
    for ch in t:
        th = "฀" <= ch <= "๿"
        w = _font("th" if th else "lat", bold).getlength(ch) / 100 * pt / 72
        tot += w * (1.05 if th else 0.93)
    return tot
def fit_one_line(t, w, pt, min_pt=11, bold=True):
    while pt > min_pt and width_in(t, pt, bold) > w * 0.88:
        pt -= 0.5
    return pt

def new_pres():
    p = Presentation(); p.slide_width = Inches(SW); p.slide_height = Inches(SH); return p
def blank(p): return p.slides.add_slide(p.slide_layouts[6])

def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, radius=None, shadow=False):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    if shadow:
        spPr = sp._element.spPr
        ef = spPr.makeelement(qn('a:effectLst'), {}); spPr.append(ef)
        sh = ef.makeelement(qn('a:outerShdw'), {'blurRad': '50000', 'dist': '20000', 'dir': '5400000', 'rotWithShape': '0'}); ef.append(sh)
        c = sh.makeelement(qn('a:srgbClr'), {'val': '8A99AD'}); sh.append(c)
        a = c.makeelement(qn('a:alpha'), {'val': '30000'}); c.append(a)
    return sp

def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True, sa=2, line_spacing=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.03); tf.margin_right = Inches(0.03); tf.margin_top = Inches(0.01); tf.margin_bottom = Inches(0.01)
    if isinstance(runs, str): runs = [[(runs, {})]]
    if isinstance(runs[0], tuple): runs = [runs]
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = align; p.space_after = Pt(sa); p.space_before = Pt(0)
        if line_spacing: p.line_spacing = line_spacing
        if isinstance(para, tuple): para = [para]
        for t, o in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(o.get("sz", 14)); r.font.bold = o.get("b", False); r.font.italic = o.get("i", False)
            r.font.name = o.get("f", FONT); r.font.color.rgb = o.get("c", GREEN)
    return tb

def logo(s, x, y, size):
    s.shapes.add_picture(os.path.join(ASSET, LOGO), Inches(x), Inches(y), Inches(size), Inches(size))

def footer(s, page):
    logo(s, 0.273, 7.052, 0.26)
    text(s, 0.62, 6.948, 3.55, 0.438, [[("THAI SUMMIT GROUP", {"sz": 20, "b": True, "c": GREEN})]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 10.28, 7.12, 2.75, 0.32, [[(str(page), {"sz": 12, "c": GREEN})]], align=PP_ALIGN.RIGHT)

def head(s, title, subtitle=None):
    text(s, 0.28, 0.28, 12.5, 0.71, [[(title, {"sz": fit_one_line(title, 12.5, 36, 18), "b": True, "c": GREEN})]])
    if subtitle:
        text(s, 0.42, 1.05, 12.3, 0.44, [[(subtitle, {"sz": fit_one_line(subtitle, 12.3, 20, 11), "b": True, "c": GREEN})]])

def headline_box(s, x, y, w, label, h=0.44):
    rect(s, x, y, w, h, fill=GREEN_D)
    text(s, x, y, w, h, [[(label, {"sz": fit_one_line(label, w - 0.16, 20, 10), "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

DECL_EN = "“I will not accept, produce and deliver non-quality work”"
DECL_TH = "“ผมจะไม่รับ, ไม่ทำและไม่ส่งมอบงานที่ไม่มีคุณภาพ”"

def title_slide(p, title, subtitle, presenter, org, quality=None, decl_label="My Quality Declaration"):
    """ปก R01 — พื้นขาว · โลโก้บนกลาง · หัวเรื่องเขียว 40 · คำประกาศคุณภาพ EN+TH · แถบรูป 4 ใบท้ายสไลด์
    (`quality` คงไว้เพื่อให้สคริปต์เดิมเรียกได้ — R01 ใช้คำประกาศคุณภาพของ template แทน)"""
    s = blank(p)
    logo(s, 5.92, 0.28, 1.25)
    text(s, 0.6, 1.66, 12.13, 0.70, [[(title, {"sz": fit_one_line(title, 12.13, 40, 20), "b": True, "c": GREEN})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        text(s, 0.6, 2.40, 12.13, 0.36, [[(subtitle, {"sz": fit_one_line(subtitle, 12.13, 18, 11), "b": True, "c": ORANGE})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.0, 2.92, 9.33, 0.95, [[(presenter, {"sz": 16, "c": GREEN})], [(org, {"sz": 16, "c": GREEN})]], align=PP_ALIGN.CENTER, sa=4)
    text(s, 3.5, 4.33, 6.33, 0.4, [[(decl_label, {"sz": 18, "b": True, "i": True, "c": GREEN})]], align=PP_ALIGN.CENTER)
    text(s, 2.10, 4.85, 9.13, 0.4, [[(DECL_EN, {"sz": 17, "c": GREEN})]], align=PP_ALIGN.CENTER)
    text(s, 2.10, 5.30, 9.13, 0.4, [[(DECL_TH, {"sz": 17, "c": GREEN})]], align=PP_ALIGN.CENTER)
    for f, (x, y, w, h) in zip(["title_a.jpg", "title_b.jpg", "title_c.jpg", "title_d.jpg"],
                               [(4.85, 6.10, 2.10, 1.27), (7.08, 6.10, 2.03, 1.26), (9.23, 6.07, 2.02, 1.30), (11.38, 6.10, 2.02, 1.27)]):
        _cover_pic(s, os.path.join(ASSET, f), x, y, w, min(h, SH - y))
    return s

def _cover_pic(s, path, x, y, w, h):
    """ใส่รูปแบบ cover (ครอปให้เต็มกรอบโดยไม่บิดสัดส่วน)"""
    from PIL import Image
    iw, ih = Image.open(path).size
    pic = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    r_img, r_box = iw / ih, w / h
    if r_img > r_box:   # รูปกว้างกว่า → ตัดซ้ายขวา
        cut = (1 - r_box / r_img) / 2; pic.crop_left = cut; pic.crop_right = cut
    else:
        cut = (1 - r_img / r_box) / 2; pic.crop_top = cut; pic.crop_bottom = cut
    return pic

def section_divider(p, photo, title_lines):
    """divider R01 — รูปเต็มฝั่งขวา + ระนาบขาวขอบเฉียง + หัวข้อเขียวฝั่งซ้าย (ไม่มี footer)"""
    s = blank(p)
    _cover_pic(s, os.path.join(ASSET, photo), 5.92, 0, 7.41, 7.5)
    rect(s, -0.05, 0, 6.1, 7.5, fill=WHITE)
    tri = rect(s, 6.05, 0, 1.31, 7.5, fill=WHITE, shape=MSO_SHAPE.RIGHT_TRIANGLE)
    tri._element.spPr.find(qn('a:xfrm')).set('flipV', '1')
    runs = [[(t, {"sz": fit_one_line(t, 5.6, 40, 20), "b": True, "c": GREEN})] for t in title_lines]
    text(s, 0.3, 2.6, 5.9, 1.9, runs, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sa=4)
    return s

def agenda_slide(p, page, title, subtitle, items):
    """Agenda R01 — ลิสต์เลขสีเขียว ไม่มีวงกลมส้ม · items = [(หัวข้อ, บรรทัดรอง)]"""
    s = blank(p)
    head(s, title, subtitle)
    yy = 1.93
    for i, (t1, t2) in enumerate(items):
        text(s, 1.56, yy, 11.0, 0.42, [[(f"{i + 1}.   {t1}", {"sz": 18, "b": True, "c": GREEN})]])
        if t2:
            text(s, 2.05, yy + 0.42, 10.5, 0.36, [[(t2, {"sz": 13, "c": GREY})]])
        yy += 0.92
    footer(s, page)
    return s

def closing_slide(p, page, motto="Before We Build Parts, We Build People", lines=()):
    """ปิด R01 — พื้นขาว คำขวัญเขียว 36 Bold + footer"""
    s = blank(p)
    text(s, 0.9, 3.04, 11.54, 0.71, [[(motto, {"sz": fit_one_line(motto, 11.54, 36, 20), "b": True, "c": GREEN})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for i, (t, c) in enumerate(lines):
        text(s, 0.9, 3.95 + i * 0.6, 11.54, 0.5, [[(t, {"sz": fit_one_line(t, 11.54, 22, 12), "b": True, "c": c})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, page)
    return s

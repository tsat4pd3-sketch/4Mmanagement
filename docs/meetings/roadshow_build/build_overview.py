# -*- coding: utf-8 -*-
"""ESM Program Overview deck — TSG theme **R01** · ภาษาอังกฤษ + ภาษาไทย (2026-10-12)
"แผนกไหนใช้อะไรได้บ้าง · ฟีเจอร์ที่เชื่อมกัน · เน้น realtime / ลดภาระ / ใช้ง่าย / ลดต้นทุน"
Run:  python3 build_overview.py
   →  ../TSAT_ESM_Program_Overview_2026-10_EN.pptx  +  ../TSAT_ESM_Program_Overview_2026-10_TH.pptx
Sources: roadshow-dept-stories-2026-10-06.md (ฟีเจอร์ + การส่งต่องาน อ้างเอกสารโมดูล)
         roadshow-metrics-2026-10-06.md (ตัวเลขทุกตัวบนสไลด์ · วัด 06 ต.ค. 2026)
กฎ: ข้อความทุกชิ้นเขียนคู่ (EN, TH) ในไฟล์นี้ที่เดียว · โน้ตผู้พูด = ไทยทั้ง 2 ฉบับ
    ห้ามใส่ตัวเลขเงินบาท (ยังไม่ได้วัด — ต้องคำนวณกับฝ่ายบัญชี)
    สีตาม R01: เขียว + ส้ม + สีที่ 3 (amber) เท่านั้น (กฎ check list ข้อ 4)
"""
import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import tsg_theme as T
from tsg_theme import GREEN, GREEN_D, GREEN_TINT, BORDER, ORANGE, AMBER, GREY, WHITE

HERE = os.path.dirname(os.path.abspath(__file__))
LANG = 0  # 0 = EN · 1 = TH (ตั้งใน build())


def t(x):
    """ข้อความคู่ (en, th) → ตามภาษาที่กำลังสร้าง · สตริงเดี่ยว = ใช้ทั้ง 2 ภาษา"""
    return x[LANG] if isinstance(x, tuple) else x


MEASURED = ("Measured from ESM database · 06 Oct 2026", "วัดจากฐานข้อมูล ESM · 6 ต.ค. 2569")
PILLARS = [(("REALTIME", "เห็นทันที"), ORANGE), (("LESS WORK", "ลดภาระงาน"), GREEN),
           (("EASY TO USE", "ใช้งานง่าย"), GREEN_D), (("LOWER COST", "ลดต้นทุน"), AMBER)]

# ═══════════════════════════ เนื้อหา (EN, TH) ═══════════════════════════
COVER = dict(
    title=("ESM PROGRAM OVERVIEW", "ภาพรวมโปรแกรม ESM"),
    sub=("Enterprise Shopfloor Management — one connected system for every department",
         "Enterprise Shopfloor Management — ระบบเดียวที่เชื่อมทุกแผนกในโรงงาน"),
    who=("Who uses what · how the work connects · why it matters",
         "แผนกไหนใช้อะไร · งานเชื่อมกันอย่างไร · ได้ประโยชน์อะไร"),
    org=("TSAT · Thai Summit Group · October 2026", "TSAT · Thai Summit Group · ตุลาคม 2569"),
)
AGENDA = [
    (("ESM at a glance", "ESM ในหน้าเดียว"), ("One login · every department · PC, phone and TV", "บัญชีเดียว · ทุกแผนก · ใช้ได้ทั้ง PC มือถือ และจอ TV")),
    (("Who uses what", "แผนกไหนใช้อะไร"), ("Departments and their screens on one page", "ทุกแผนกกับจอที่ใช้ในหน้าเดียว")),
    (("How the work connects", "งานเชื่อมกันอย่างไร"), ("Two end-to-end flows across departments", "2 สายงานที่วิ่งข้ามแผนกตั้งแต่ต้นจนจบ")),
    (("Department by department", "ลงรายละเอียดทีละแผนก"), ("Features, key screens and hand-offs", "ฟีเจอร์ · จอหลัก · การส่งต่องาน")),
    (("Why it matters", "ได้ประโยชน์อะไร"), ("Real-time · Less work · Easy to use · Lower cost", "เห็นทันที · ลดภาระงาน · ใช้งานง่าย · ลดต้นทุน")),
]
GLANCE = [
    ["Live OEE during the shift", "Andon on factory map & TV", "Calls go to the right team's phone", "Escalates by itself if unanswered"],
    ["No re-typing paper into Excel", "Official FM- forms printed from data", "OEE, PM due dates, kanban calculated", "System opens follow-up orders"],
    ["Scan QR on machine / jig / part", "Thai screens, one login for all", "Works on phone, PC and factory TV", "'My queue' shows what waits for me"],
    ["Built in-house — no per-user licence", "Uses PCs, phones, TVs we already have", "~1,300 paper forms / month now digital", "Less machine waiting time for repair"],
]
GLANCE_TH = [
    ["OEE สดระหว่างกะ", "สถานะไลน์บนผังโรงงานและจอ TV", "เรียกช่างเด้งถึงมือถือทีมที่ถูกต้อง", "ไม่มีคนรับ ระบบเตือนต่อเอง"],
    ["ไม่ต้องคีย์กระดาษลง Excel ซ้ำ", "พิมพ์ฟอร์ม FM- ทางการจากข้อมูล", "คำนวณ OEE วัน PM คัมบังให้เอง", "ระบบเปิดใบงานต่อเนื่องให้เอง"],
    ["สแกน QR ที่เครื่อง จิ๊ก ชิ้นงาน", "จอภาษาไทย บัญชีเดียวทุกแผนก", "ใช้ได้ทั้งมือถือ PC จอ TV", "'คิวงานของฉัน' บอกงานที่รอเรา"],
    ["พัฒนาเอง ไม่มีค่าไลเซนส์รายคน", "ใช้ PC มือถือ จอ TV ที่มีอยู่แล้ว", "กระดาษ ~1,300 ใบ/เดือน เข้าระบบ", "เครื่องรอซ่อมสั้นลง"],
]
GLANCE_STATS = [("10", ("departments on one system", "แผนกในระบบเดียว")), ("80+", ("screens & forms", "จอและฟอร์ม")),
                ("71", ("people active in last 30 days", "คนใช้งานจริงใน 30 วัน")), ("1", ("login · PC · phone · TV", "บัญชี · PC · มือถือ · TV"))]
MATRIX = [
    (("Production", "ฝ่ายผลิต"), ("Daily Report · Downtime & NG · Check-in & PPE · LPA · Scrap report · Problem report",
                                 "รายงานผลิตรายวัน · Downtime/ของเสีย · เช็คชื่อ-PPE · LPA · ใบรายงานของเสีย · ใบรายงานปัญหา")),
    (("Maintenance", "ซ่อมบำรุง"), ("Repair orders (MO) · Andon board · PM plans & checks · Spare parts · KPI & QC 7 tools",
                                  "ใบแจ้งซ่อม MO · บอร์ด Andon ช่าง · แผน/ใบตรวจ PM · อะไหล่ · KPI ช่าง และ QC 7 Tools")),
    (("JIG / DIE", "JIG / DIE"), ("Jig & die registry · Shim record · Die status · Jig inspection · Setup-time rules",
                                  "ทะเบียนจิ๊ก/แม่พิมพ์ · บันทึกชิม · สถานะแม่พิมพ์ · ใบตรวจจิ๊ก · กฎเวลาเปลี่ยนรุ่น")),
    (("Quality (QA)", "ประกันคุณภาพ (QA)"), ("Inspection sheets · Yellow / red bins · 4M approval · CQI-15 · Test-part requisition",
                                           "ใบตรวจ · ถังเหลือง/ถังแดง · อนุมัติ 4M · CQI-15 · ใบเบิกชิ้นทดสอบ")),
    (("Store", "สโตร์"), ("Store board (WIP refill · raw mat · purchase) · Stock receipts · Store monitor · BOM auto-request",
                          "บอร์ดสโตร์ (เติม WIP · วัตถุดิบ · จัดซื้อ) · รับเข้าคลัง · เฝ้าระวังสต็อก · ขอพาร์ทลูกอัตโนมัติจาก BOM")),
    (("Logistic", "Logistic / จัดส่ง"), ("Customer demand & shipping · EDI from e-mail · FG stock · In-plant transport",
                                        "ความต้องการลูกค้าและการส่ง · EDI จากเมล · สต็อก FG · ขนส่งในโรงงาน")),
    (("Planning / Sales", "วางแผน / ขาย"), ("Production plan · Monitoring boards · Kanban calculation · Customer part matching",
                                          "แผนการผลิต · บอร์ด Monitoring · คำนวณคัมบัง · จับคู่เลขพาร์ทลูกค้า")),
    (("Engineering (PE)", "วิศวกรรม (PE)"), ("Process flow · PFMEA · Control plan · NPI / APQP / PPAP · 4M change",
                                           "Process Flow · PFMEA · Control Plan · NPI / APQP / PPAP · เปลี่ยนแปลง 4M")),
    (("HR / Admin", "HR / ธุรการ"), ("OJT training (e-sign) · Skills matrix · Workforce insight · Company calendar · OT",
                                   "อบรม OJT (เซ็นบนจอ) · ตารางทักษะ · กำลังคน · ปฏิทินบริษัท · OT")),
    (("Management", "ผู้บริหาร"), ("OBEYA KPI & SQDCM · Factory map · TV dashboards · Order trace · Action board",
                                 "OBEYA KPI และ SQDCM · ผังโรงงาน · จอ TV · สอบกลับใบผลิต · บอร์ดติดตาม Action")),
]
FLOW1 = (("①  A machine stops", "①  เครื่องจักรหยุด"), [
    (("Line records stop", "ไลน์บันทึกเครื่องหยุด"), ("Daily Report · 1 tap", "กดครั้งเดียว")),
    (("Calls the right team", "เรียกทีมช่างที่ถูกต้อง"), ("JIG / DIE / MTN phone", "เด้งมือถือทีม JIG/DIE/MTN")),
    (("Repair order opens", "เปิดใบแจ้งซ่อม"), ("MO from the downtime", "ต่อจาก downtime ทันที")),
    (("Repair → QA check", "ซ่อม → QA ตรวจ"), ("Signed on screen", "เซ็นบนจอ")),
    (("OEE updates", "OEE อัปเดต"), ("Shift + factory map", "กะนั้น + ผังโรงงาน")),
    (("KPI on OBEYA", "ขึ้น KPI บน OBEYA"), ("Management sees it", "ผู้บริหารเห็นทันที")),
])
FLOW2 = (("②  A customer order arrives", "②  ลูกค้าสั่งของเข้ามา"), [
    (("EDI e-mail in", "เมล EDI เข้า"), ("Auto queue · 15 min", "เข้าคิวเองทุก 15 นาที")),
    (("Plan & monitor", "วางแผน / ติดตาม"), ("Shifts / OT needed", "ต้องเปิดกี่กะ / OT")),
    (("Line produces", "ไลน์ผลิต"), ("Scan kanban card", "สแกนบัตรคัมบัง")),
    (("Close order", "ปิดใบผลิต"), ("BOM requests parts", "ระบบขอพาร์ทลูกจาก BOM")),
    (("Store & FG", "สโตร์ / คลัง FG"), ("Count, then receive", "นับก่อนรับเข้า")),
    (("Ship", "ส่งของ"), ("Cuts stock · red if short", "ตัดสต็อก · ไม่พอขึ้นแดง")),
])
TRACE = (("One order number traces everything:  ", "เลขใบผลิตเดียว สอบกลับได้ทั้งเส้น:  "),
         ("which line, shift, people, machine stops, defects, repairs and the PFMEA line behind it — on one screen (Order Trace).",
          "ไลน์ไหน กะไหน ใครทำ เครื่องหยุดกี่ครั้ง ของเสีย การซ่อม จนถึงบรรทัด PFMEA ที่เกี่ยวข้อง — บนจอเดียว (Order Trace)"))

# แต่ละแผนก: title, sub, features[], screens[], links[(who, what)], pillars[4], notes
DEPTS = [
    dict(title=("Production", "ฝ่ายผลิต"),
         sub=("Record once at the line — the rest of the factory sees it", "บันทึกครั้งเดียวที่หน้างาน แล้วทั้งโรงงานเห็นทันที"),
         feat=[("Open the shift and scan the kanban card — no order typing", "เปิดกะแล้วสแกนบัตรคัมบัง ไม่ต้องพิมพ์เลขใบผลิต"),
               ("Log stops and defects with reason; live OEE (A × P × Q) during the shift", "ลงเครื่องหยุดและของเสียพร้อมสาเหตุ เห็น OEE (A × P × Q) สดระหว่างกะ"),
               ("Call maintenance by team with one button", "เรียกช่างแยกทีมได้ด้วยปุ่มเดียว"),
               ("Daily check-in, PPE and skills fit for each workstation", "เช็คชื่อ PPE และความเหมาะสมของทักษะต่อจุดงานทุกวัน"),
               ("LPA audits and scrap / problem reports printed as official FM- forms", "LPA ใบรายงานของเสีย และใบรายงานปัญหา พิมพ์เป็นฟอร์ม FM- ทางการ"),
               ("Shift approval locks the numbers — no editing afterwards", "อนุมัติปิดกะแล้วตัวเลขถูกล็อก แก้ย้อนหลังไม่ได้")],
         scr=[("Daily Report", "รายงานผลิตรายวัน"), ("Check-in & PPE", "เช็คชื่อ / PPE"), ("Layer Process Audit", "LPA"),
              ("Scrap Report", "ใบรายงานของเสีย"), ("Problem report & bins", "ใบรายงานปัญหา / ถัง"), ("Factory map / TV", "ผังโรงงาน / จอ TV")],
         links=[("MTN", ("Calls + repair order opened from the stop", "เรียกช่าง + เปิดใบซ่อมต่อจากเครื่องหยุด")),
                ("QA", ("Suspect parts to yellow / red bins", "ของสงสัยเข้าถังเหลือง/แดง")),
                (("Store", "สโตร์"), ("Closing an order requests child parts", "ปิดใบผลิตแล้วขอพาร์ทลูกให้เอง")),
                (("Mgmt", "ผู้บริหาร"), ("OEE flows to OBEYA KPI", "OEE ไหลเข้า KPI บน OBEYA"))],
         pil=[("OEE known during the shift, not at month-end", "รู้ OEE ระหว่างกะ ไม่ต้องรอสิ้นเดือน"),
              ("No hand-written daily report or re-typing", "เลิกเขียนรายงานมือและคีย์ซ้ำ"),
              ("Scan instead of choosing from long lists", "สแกนแทนการเลือกจากลิสต์ยาว"),
              ("Faster calls = less line waiting", "เรียกช่างเร็ว ไลน์รอน้อยลง")],
         th="ผลิตคือต้นทาง: บันทึกครั้งเดียวที่หน้างาน แล้วข้อมูลวิ่งไปทุกแผนก\n"
            "ตัวเลข ก.ย.: รายงานกะในระบบ 505 ใบ · downtime 3,792 รายการ · เช็คชื่อ 3,489 ครั้ง · กะที่ OEE พร้อมตอนปิดกะ 70% → 96% (ก.ค. → ก.ย.)\n"
            "เดโม: /daily-report เปิดกะ → ลง downtime → กดเรียกช่าง"),
    dict(title=("Maintenance (MTN)", "ซ่อมบำรุง (MTN)"),
         sub=("From a phone call to a tracked repair order", "จากการโทรตาม สู่ใบแจ้งซ่อมที่ติดตามได้ทุกขั้น"),
         feat=[("Calls arrive on the right team's phone and Andon board — with siren", "การเรียกเด้งถึงมือถือทีมที่ถูกต้องและบอร์ด Andon พร้อมเสียงเตือน"),
               ("Repair order (MO) opens straight from the machine stop", "เปิดใบแจ้งซ่อม MO ต่อจากเครื่องหยุดได้ทันที"),
               ("Repairs start at once; approvals are signed afterwards on screen", "งานซ่อมเริ่มได้เลย เซ็นอนุมัติตามหลังบนจอ"),
               ("PM plans, due dates and check sheets — overdue shows red", "แผน PM วันครบกำหนด และใบตรวจ เลยกำหนดขึ้นแดง"),
               ("Spare parts used per MO; rank A / B / C automatically", "อะไหล่ที่ใช้ต่อใบ MO จัดอันดับ A/B/C ให้เอง"),
               ("Technician KPI, MTTR / MTBF and QC 7 tools built in", "KPI ช่าง MTTR / MTBF และ QC 7 Tools ในตัว")],
         scr=[("Repair orders (MO)", "ใบแจ้งซ่อม MO"), ("MTN Andon board", "บอร์ด Andon ช่าง"), ("PM forecast & checks", "PM ล่วงหน้า / ใบตรวจ"),
              ("PM coordination", "แผนประสานงาน PM"), ("Equipment & spares", "ทะเบียนอุปกรณ์ / อะไหล่"), ("MTN analysis / QC7", "วิเคราะห์งานซ่อม / QC7")],
         links=[(("Prod", "ผลิต"), ("Receives calls; told when the line can run", "รับเรียกจากไลน์ แจ้งกลับเมื่อเดินเครื่องได้")),
                ("QA", ("Checks the repair before closing", "ตรวจงานซ่อมก่อนปิดใบ")),
                (("Store", "สโตร์"), ("Spare-part stock and usage", "สต็อกและการเบิกอะไหล่")),
                (("Plan", "วางแผน"), ("Sees PM days ahead to build buffer", "เห็นวัน PM ล่วงหน้า ผลิตเผื่อได้"))],
         pil=[("Call acknowledged P80: 529 → 30 min", "เรียกช่างถึงรับเรื่อง P80: 529 → 30 นาที"),
              ("No paper FM-MTN-006 / FM-JIG-008", "เลิกใช้ฟอร์มกระดาษ FM-MTN-006 / FM-JIG-008"),
              ("Everything on the phone at the machine", "ทำทุกอย่างบนมือถือที่หน้าเครื่อง"),
              ("Repair done P80: 7.9 h → 0.9 h", "ซ่อมเสร็จ P80: 7.9 → 0.9 ชม.")],
         th="MTN คือแผนกที่เห็นผลชัดที่สุด\n"
            "วัดจริง ส.ค. → ก.ย.: เรียกช่างถึงรับเรื่อง (P80) 529 → 30 นาที · เครื่องหยุดถึงเปิด MO (ค่ากลาง) 7.7 → 3.2 นาที · เปิด MO ถึงซ่อมเสร็จ (P80) 7.9 → 0.9 ชม.\n"
            "ใบ MO ในระบบ ส.ค. 76 → ก.ย. 496 ใบ · ใบตรวจ PM 6 → 88 · แผน PM 147 แผน เลยกำหนดวันนี้ 1\n"
            "ข้อควรระวัง: ใบค้างรอ QA ยังเยอะ (205 ใบ) — ถ้าถูกถามให้ตอบตรงๆ"),
    dict(title=("JIG & DIE", "JIG และ DIE"),
         sub=("Tooling history in one place instead of separate files", "ประวัติจิ๊กและแม่พิมพ์อยู่ที่เดียว แทนไฟล์แยกกระจาย"),
         feat=[("Registry of jigs and dies with QR labels", "ทะเบียนจิ๊กและแม่พิมพ์ พร้อมป้าย QR"),
               ("Jig inspection points and photos per checkpoint", "จุดตรวจจิ๊กพร้อมรูปทุกจุด"),
               ("Fixture shim record — customer requirement for sustainability", "บันทึกการใส่ชิม (shim) ตามข้อกำหนดลูกค้า"),
               ("Die status: ready / in repair / at the press", "สถานะแม่พิมพ์: พร้อม / ซ่อม / อยู่ที่เครื่อง"),
               ("Setup-time rules per press (die height, shut height)", "กฎเวลาเปลี่ยนรุ่นต่อเครื่อง (ความสูงแม่พิมพ์ / shut height)"),
               ("JIG / DIE team receives only its own calls", "ทีม JIG / DIE ได้รับเฉพาะงานของทีมตัวเอง")],
         scr=[("Equipment center", "ศูนย์ทะเบียนอุปกรณ์"), ("Jig registry & shim", "ทะเบียนจิ๊ก / ชิม"), ("Die maintenance", "ซ่อมบำรุงแม่พิมพ์"),
              ("PM setup (jig checks)", "ตั้งค่า PM (จุดตรวจจิ๊ก)"), ("Repair orders (MO)", "ใบแจ้งซ่อม MO")],
         links=[(("Prod", "ผลิต"), ("Calls JIG / DIE team directly", "เรียกทีม JIG / DIE ได้ตรง")),
                ("MTN", ("Same MO workflow, own team", "ใช้ขั้นตอน MO เดียวกัน แยกทีม")),
                (("Plan", "วางแผน"), ("Die setup time used in the plan", "เวลาเปลี่ยนแม่พิมพ์ใช้ในแผน")),
                ("QA", ("Checks after jig repair", "ตรวจหลังซ่อมจิ๊ก"))],
         pil=[("Only the right team is alerted", "แจ้งเตือนเฉพาะทีมที่เกี่ยวข้อง"),
              ("One registry — no separate Excel files", "ทะเบียนเดียว เลิกไฟล์ Excel แยก"),
              ("Scan the tool to open its history", "สแกนแล้วเห็นประวัติทันที"),
              ("Fewer wrong-team calls and walking", "ลดการเรียกผิดทีมและการเดินตาม")],
         th="JIG/DIE ใช้ระบบเดียวกับ MTN แต่แยกทีม — ปุ่มเรียกช่างต้องเลือกทีม เด้งเฉพาะทีมนั้น (เดิมเรียกทั้ง 13 คนทุกทีม)\n"
            "สถานะจริง: ทะเบียนพร้อม · ใบตรวจจิ๊กและบันทึกชิมยังลงน้อย — พูดเป็น 'ระบบพร้อมใช้' ไม่ใช่ 'ใช้เต็มแล้ว'"),
    dict(title=("Quality (QA)", "ประกันคุณภาพ (QA)"),
         sub=("Quality decisions recorded where the parts are", "ตัดสินคุณภาพและบันทึกผลที่หน้างาน"),
         feat=[("Inspection sheets piece by piece — out-of-spec cannot be passed", "ใบตรวจทีละชิ้น ค่าหลุดสเปกกดผ่านไม่ได้"),
               ("Yellow / red bin tags with age; QA decides scrap / repair / use-as-is", "แท็กถังเหลือง/แดงนับอายุ QA ตัดสิน ทิ้ง / ซ่อม / ใช้ได้"),
               ("4M change approval with Telegram alert at every step", "อนุมัติ 4M แจ้งเตือน Telegram ทุกขั้น"),
               ("CQI-15 welding event log", "บันทึกเหตุการณ์งานเชื่อม CQI-15"),
               ("Test-part requisition FM-STO-003 replaces hand-written slips", "ใบเบิกชิ้นทดสอบ FM-STO-003 แทนใบเขียนมือ"),
               ("Suspect parts stay out of %Q until QA decides — numbers stay honest", "ของสงสัยยังไม่นับใน %Q จนกว่า QA ตัดสิน ตัวเลขจึงตรงจริง")],
         scr=[("QA inspection", "ใบตรวจ QA"), ("Quality bins", "ถังเหลือง / แดง"), ("4M approval", "อนุมัติ 4M"),
              ("CQI-15 event log", "บันทึก CQI-15"), ("Test-part requisition", "ใบเบิกชิ้นทดสอบ")],
         links=[(("Prod", "ผลิต"), ("Bins and 4M come from the line", "ถังและใบ 4M มาจากไลน์")),
                ("MTN", ("QA checks repairs before close", "QA ตรวจงานซ่อมก่อนปิดใบ")),
                (("Store", "สโตร์"), ("Issues test parts to QA", "จ่ายชิ้นทดสอบให้ QA")),
                ("PE", ("Defects link to PFMEA lines", "ของเสียโยงถึงบรรทัด PFMEA"))],
         pil=[("Bin age alerts before tags expire", "เตือนก่อนแท็กถังหมดอายุ"),
              ("No re-writing 3 paper forms per problem", "เลิกเขียนกระดาษ 3 ใบต่อปัญหา"),
              ("Decide on screen at the bin", "ตัดสินบนจอที่หน้าถัง"),
              ("Scrap counted once, at the right cost", "นับของเสียครั้งเดียว คิดต้นทุนถูก")],
         th="QA: ระบบพร้อมครบ แต่การใช้งานยังน้อย (ผลตัดสินถัง = 0 · ใบตรวจล่าสุด ส.ค.)\n"
            "ตอนนำเสนอให้เน้นว่า 'ระบบรองรับอะไร' และยกเป็นเป้าเดือนหน้า — อย่าพูดว่าใช้เต็มแล้ว"),
    dict(title=("Store", "สโตร์"),
         sub=("From LINE chat requests to a tracked queue", "จากขอของผ่านแชท LINE สู่คิวงานที่ติดตามได้"),
         feat=[("Store board: WIP refill, child parts, raw material, purchase, FG", "บอร์ดสโตร์: เติม WIP · พาร์ทลูก · วัตถุดิบ · จัดซื้อ · FG"),
               ("Closing a production order explodes the BOM and requests parts by itself", "ปิดใบผลิตแล้วระบบกระจาย BOM ขอพาร์ทลูกให้เอง"),
               ("Scan the part (wrong part = blocked) and the delivery point QR", "สแกนชิ้นงาน (ผิดตัว = บล็อก) และ QR จุดส่ง"),
               ("Count before receiving; differences need a reason", "นับก่อนรับเข้า ยอดไม่ตรงต้องใส่เหตุผล"),
               ("Manual stock adjustments wait for approval", "ปรับยอดสต็อกเองต้องรออนุมัติ"),
               ("Store monitor: below min, over max, late rounds, short receipts", "เฝ้าระวัง: ต่ำกว่า Min · เกิน Max · รอบส่งเลยเวลา · รับไม่ครบ")],
         scr=[("Heijunka / store board", "บอร์ดสโตร์ / Heijunka"), ("Line stock & receipts", "สต็อกไลน์ / รับเข้า"), ("Store monitor", "เฝ้าระวังสต็อก"),
              ("Flow tower", "Flow Tower"), ("Material requests", "ใบเบิกวัสดุ")],
         links=[(("Prod", "ผลิต"), ("Line requests refill; closes orders", "ไลน์ขอเติมของ / ปิดใบผลิต")),
                (("Purch.", "จัดซื้อ"), ("Shortage creates purchase requests", "ของขาดเปิดใบขอซื้อให้เอง")),
                ("QA", ("Issues test parts with approval", "จ่ายชิ้นทดสอบตามใบอนุมัติ")),
                (("Logist.", "จัดส่ง"), ("FG receipts feed shipping", "รับ FG เข้าคลังเพื่อส่ง"))],
         pil=[("Late rounds and short stock turn red", "รอบส่งเลยเวลาหรือของขาดขึ้นแดง"),
              ("No reading requests from chat groups", "เลิกไล่อ่านคำขอจากแชทกลุ่ม"),
              ("Scan-to-pick, scan-to-deliver", "สแกนหยิบ สแกนส่ง"),
              ("Less excess stock and fewer stock-outs", "สต็อกเกินลดลง ของขาดน้อยลง")],
         th="สโตร์: เดิมรับคำขอผ่าน LINE กลุ่ม (ทีมสโตร์เล่าเองว่า 'มั่วมาก')\n"
            "ตอนนี้ขาที่ระบบทำเองครบ (ระเบิด BOM ออกใบเอง) · ขาที่ต้องมีคนกด (รับใบ/บันทึกจ่าย) ยังต้องฝึกเพิ่ม\n"
            "รับเข้าคลังหลังปิดใบผลิต ค่ากลางยัง ~73 ชม. — ถ้าถูกถาม ตอบว่าโหมดนับก่อนรับเพิ่งเริ่ม 02/10"),
    dict(title=("Logistic & Delivery", "Logistic และการจัดส่ง"),
         sub=("Customer orders arrive by themselves; shipping cuts stock", "ยอดสั่งลูกค้าเข้าระบบเอง ส่งของแล้วตัดสต็อกทันที"),
         feat=[("EDI 830 / 862 attachments pulled from e-mail into a queue", "ไฟล์แนบ EDI 830 / 862 ดึงจากเมลเข้าคิวเอง"),
               ("Customer call-off (e-SMART) overrides the plan as final quantity", "ยอดเรียกของลูกค้า (e-SMART) เป็นยอดสุดท้ายแทนแผน"),
               ("Shipping countdown in 4 phases; stock enough / short / none in colour", "นับถอยหลังการส่ง 4 เฟส สีบอก สต็อกพอ / ขาด / ไม่มี"),
               ("'Shipped' cuts FG stock at once — failure turns red, never silent", "กด 'ส่งแล้ว' ตัดสต็อก FG ทันที ตัดไม่ได้ขึ้นแดง"),
               ("FG receipts from the line: count, then receive", "รับ FG จากไลน์: นับก่อนแล้วค่อยรับ"),
               ("In-plant transport: drivers, shortest stop order, simulation", "ขนส่งในโรงงาน: มอบหมายคนขับ จัดลำดับจุดจอดสั้นสุด จำลองเส้นทาง")],
         scr=[("Customer demand", "ความต้องการลูกค้า"), ("Planner & Sales (EDI)", "Planner & Sales (EDI)"), ("FG stock", "สต็อก FG"),
              ("Transport", "ขนส่งในโรงงาน"), ("Rundown stock", "Rundown stock")],
         links=[(("Sales", "ขาย"), ("EDI and e-SMART feed demand", "EDI และ e-SMART ส่งยอดสั่ง")),
                (("Prod", "ผลิต"), ("'No stock — must produce!' alert", "เตือน 'ไม่มีสต็อก ต้องผลิต!'")),
                (("Store", "สโตร์"), ("FG receipts from closed orders", "รับ FG จากใบที่ปิดแล้ว")),
                ("MTN", ("Which shipments a breakdown hits", "เครื่องเสียกระทบรอบส่งไหน"))],
         pil=[("Late shipments show red the day before", "รอบส่งเสี่ยงขึ้นแดงล่วงหน้า 1 วัน"),
              ("No saving and dragging e-mail files", "เลิกเซฟไฟล์แนบแล้วลากเข้าเอง"),
              ("One screen per shipping round", "จอเดียวต่อรอบส่ง"),
              ("Fewer premium freights from surprises", "ลดการส่งด่วนพิเศษจากเรื่องไม่คาดคิด")],
         th="Logistic: เมล EDI เข้าคิวเองทุก 15 นาที (เดิมต้องเซฟไฟล์แนบแล้วลากเข้าเอง)\n"
            "ตัวเลขส่งตรงเวลา: ส.ค. 91% · ก.ย. 72% · ต.ค. 87% (6 วันแรก) — ผู้นำเสนอต้องเตรียมเหตุผลที่ ก.ย. ตก\n"
            "จุดที่ต้องปิด: ยังใช้บัญชีร่วม — ควรให้แต่ละคนมีบัญชีของตัวเอง"),
    dict(title=("Planning & Sales", "วางแผนและขาย"),
         sub=("Plans from real capacity, not from Excel", "วางแผนจากกำลังผลิตจริง ไม่ใช่จาก Excel"),
         feat=[("Production plan from real output (median of last 60 days)", "แผนผลิตจากยอดจริง (ค่ากลาง 60 วันล่าสุด)"),
               ("Tells how many shifts, which days need OT, night shift or holiday work", "บอกว่าต้องเปิดกี่กะ วันไหนต้อง OT กะดึก หรือทำวันหยุด"),
               ("Monitoring boards replace the 13-sheet daily Excel", "บอร์ด Monitoring แทน Excel 13 ชีทที่ทำทุกวัน"),
               ("Kanban calculation from forecast with the company calendar", "คำนวณคัมบังจาก forecast ตามปฏิทินบริษัท"),
               ("Customer part numbers matched to MAT; unmatched ones are flagged", "จับคู่เลขพาร์ทลูกค้ากับ MAT จับไม่ได้ระบบเตือน"),
               ("Plan vs actual per part — over and off-plan shown, never hidden", "แผนเทียบผลจริงรายพาร์ท ผลิตเกิน/นอกแผนโชว์ ไม่ซ่อน")],
         scr=[("Production plan", "แผนการผลิต"), ("Monitoring boards", "บอร์ด Monitoring"), ("Kanban calculation", "คำนวณคัมบัง"),
              ("Planner & Sales", "Planner & Sales"), ("Company calendar", "ปฏิทินบริษัท")],
         links=[(("Sales", "ขาย"), ("Orders and forecast in", "ยอดสั่งและ forecast เข้า")),
                (("Prod", "ผลิต"), ("Plan = frame; the line confirms", "แผน = กรอบ ไลน์เป็นคนยืนยันยอด")),
                (("Store", "สโตร์"), ("Kanban and lot size", "จำนวนคัมบังและขนาดล็อต")),
                ("MTN", ("PM days shown before planning", "เห็นวัน PM ก่อนวางแผน"))],
         pil=[("Plan vs actual updates as the line closes", "แผนเทียบจริงอัปเดตทันทีที่ไลน์ปิดใบ"),
              ("No daily 13-sheet Excel", "เลิกทำ Excel 13 ชีททุกวัน"),
              ("Same screens as the Excel the team knows", "หน้าตาเหมือน Excel ที่ทีมคุ้นเคย"),
              ("Right shifts and OT — not too many", "เปิดกะและ OT พอดี ไม่เกินจำเป็น")],
         th="Planning: เป้าคือเลิกทำ Excel Monitoring 13 ชีททุกวัน — จอ /monitoring หน้าตาเดียวกับ Excel ที่ทีมใช้\n"
            "สถานะจริง: ระบบพร้อม · การใช้แผนผลิต/คำนวณคัมบังยังน้อย — ยกเป็นเป้าเดือนหน้า"),
    dict(title=("Engineering (PE) & NPI", "วิศวกรรม (PE) และ NPI"),
         sub=("Process documents as one living data set", "เอกสารกระบวนการเป็นข้อมูลชุดเดียวที่อัปเดตตลอด"),
         feat=[("Process flow, PFMEA and control plan from one data set (by OP)", "Process Flow · PFMEA · Control Plan จากข้อมูลชุดเดียว (ตาม OP)"),
               ("Import existing Excel files; export in the original layout", "นำเข้าไฟล์ Excel เดิมได้ ส่งออกตามรูปแบบเดิม"),
               ("Central PFMEA library — system proposes, engineer decides", "คลัง PFMEA กลาง ระบบเสนอ วิศวกรตัดสิน"),
               ("Revisions linked to the claim, defect or repair that caused them", "ทุก revision โยงถึงเคลม ของเสีย หรืองานซ่อมที่เป็นต้นเหตุ"),
               ("NPI: APQP phases, PPAP elements, drawing revisions, tooling plan", "NPI: เฟส APQP · PPAP · revision แบบ · แผน tooling"),
               ("4M Method change linked to the engineering change", "เปลี่ยน 4M ด้าน Method โยงกับการเปลี่ยนแปลงทางวิศวกรรม")],
         scr=[("PE documents (PFC / PFMEA / CP)", "เอกสาร PE (PFC / PFMEA / CP)"), ("PFMEA master", "คลัง PFMEA กลาง"),
              ("NPI / APQP", "NPI / APQP"), ("4M change", "เปลี่ยนแปลง 4M"), ("Order trace", "สอบกลับใบผลิต")],
         links=[("QA", ("Defects and 8D point to PFMEA lines", "ของเสียและ 8D ชี้ถึงบรรทัด PFMEA")),
                (("Prod", "ผลิต"), ("Control plan at the line", "Control Plan ใช้ที่หน้างาน")),
                ("MTN", ("Repairs can trigger a revision", "งานซ่อมเป็นเหตุให้แก้ revision ได้")),
                (("Mgmt", "ผู้บริหาร"), ("New-part status per customer", "สถานะพาร์ทใหม่รายลูกค้า"))],
         pil=[("Change once — flow, FMEA and CP stay in sync", "แก้ที่เดียว Flow / FMEA / CP ตรงกันเสมอ"),
              ("No 3 separate Excel files per part", "เลิกไฟล์ Excel 3 ไฟล์ต่อพาร์ท"),
              ("See every revision and why", "เห็นทุก revision และเหตุผล"),
              ("Fewer repeat claims via yokoten", "ขยายผลข้ามพาร์ท ลดเคลมซ้ำ")],
         th="PE: 3 ไฟล์ Excel (PFC/PFMEA/CP) กลายเป็นข้อมูลชุดเดียว มองได้ 3 มุม\n"
            "พาร์ทต้นแบบ P703 (สาย 060/061) ลงข้อมูลเต็มใบแล้ว ใช้เดโมได้\n"
            "สถานะจริง: เอกสารลงล่าสุด ส.ค. · คลัง PFMEA กลางรอ PE ยืนยัน 43 กระบวนการ"),
    dict(title=("HR & Admin", "HR และธุรการ"),
         sub=("People, skills and the calendar every plan depends on", "คน ทักษะ และปฏิทินที่ทุกแผนต้องพึ่ง"),
         feat=[("OJT training FM-HRM-004 with on-screen signatures", "ใบอบรม OJT FM-HRM-004 เซ็นบนจอ"),
               ("Skills matrix and individual evaluation forms printed from data", "ตารางทักษะและใบประเมินรายบุคคลพิมพ์จากข้อมูล"),
               ("Daily check-in, OT and holiday transport booking", "เช็คชื่อรายวัน OT และจองรถวันหยุด"),
               ("Company calendar drives OT, kanban, plans, PM and LPA", "ปฏิทินบริษัทคุม OT คัมบัง แผนผลิต PM และ LPA"),
               ("Workforce insight: daily manpower, station changes, turnover", "ข้อมูลกำลังคน: รายวัน การย้ายจุดงาน และ turnover"),
               ("Monthly KPI packs exported in the HRM form layout", "ส่งออกชุด KPI รายเดือนตามฟอร์ม HRM")],
         scr=[("OJT training", "อบรม OJT"), ("Skills matrix / operator", "ตารางทักษะ / พนักงาน"), ("Workforce insight", "ข้อมูลกำลังคน"),
              ("Company calendar", "ปฏิทินบริษัท"), ("Manpower board", "บอร์ดกำลังคน")],
         links=[(("Prod", "ผลิต"), ("Skills fit for each workstation", "ทักษะเหมาะกับจุดงาน")),
                (("All", "ทุกแผนก"), ("Calendar errors reach every plan", "ปฏิทินผิด กระทบทุกแผน")),
                (("Mgmt", "ผู้บริหาร"), ("Manpower and turnover", "กำลังคนและ turnover")),
                ("QSM", ("KPI packs in the official form", "ชุด KPI ตามฟอร์มทางการ"))],
         pil=[("Manpower per line seen today", "เห็นกำลังคนรายไลน์วันนี้"),
              ("No printing and signing paper packs", "เลิกพิมพ์และเซ็นกระดาษเป็นปึก"),
              ("Sign on the screen", "เซ็นบนจอ"),
              ("Less overtime planned by mistake", "ลด OT ที่วางผิด")],
         th="HR/ธุรการ: ปฏิทินบริษัทสำคัญมาก — ปฏิทินผิด = คัมบัง/แผนผลิต/PM/LPA ผิดหมดเงียบๆ (เคยเกิดจริง 28–29 ก.ย.)\n"
            "ใบ OJT ในระบบ: ส.ค. 12 · ก.ย. 6 ใบ"),
    dict(title=("Management & OBEYA", "ผู้บริหารและ OBEYA"),
         sub=("The whole factory on one wall — and on your phone", "เห็นทั้งโรงงานบนผนังเดียว และบนมือถือ"),
         feat=[("OBEYA KPI board and SQDCM — actual numbers come from the floor", "บอร์ด OBEYA KPI และ SQDCM ตัวเลขจริงมาจากหน้างาน"),
               ("Factory map: every line's Andon colour on one screen", "ผังโรงงาน: สถานะ Andon ทุกไลน์บนจอเดียว"),
               ("TV dashboards per department, controlled from a phone", "จอ TV รายแผนก ควบคุมจากมือถือ"),
               ("Order trace: one number shows the whole history", "สอบกลับใบผลิต: เลขเดียวเห็นประวัติทั้งหมด"),
               ("Action board: off-target → owner → due date", "บอร์ด Action: หลุดเป้า → ผู้รับผิดชอบ → กำหนดเสร็จ"),
               ("'Program update' shows what changed each week", "'อัพเดทโปรแกรม' บอกว่าแต่ละสัปดาห์เปลี่ยนอะไร")],
         scr=[("OBEYA", "OBEYA"), ("Factory map", "ผังโรงงาน"), ("TV / department dashboards", "จอ TV / Dashboard แผนก"),
              ("Order trace", "สอบกลับใบผลิต"), ("Program update", "อัพเดทโปรแกรม")],
         links=[(("All", "ทุกแผนก"), ("Every KPI comes from department records", "KPI ทุกตัวมาจากบันทึกของแผนก")),
                (("Prod", "ผลิต"), ("Andon and OEE per line", "Andon และ OEE รายไลน์")),
                ("MTN", ("Downtime cost and MTTR", "ต้นทุนเครื่องหยุดและ MTTR")),
                (("Plan", "วางแผน"), ("Plan vs actual per part", "แผนเทียบจริงรายพาร์ท"))],
         pil=[("KPI updated as the floor records", "KPI อัปเดตตามที่หน้างานบันทึก"),
              ("No printing KPI sheets for the wall", "เลิกพิมพ์ KPI แปะผนัง"),
              ("One screen answers 'where are we?'", "จอเดียวตอบ 'ตอนนี้เราอยู่ตรงไหน'"),
              ("Decide on facts, faster", "ตัดสินใจจากข้อเท็จจริงได้เร็วขึ้น")],
         th="ผู้บริหาร: บอร์ด OBEYA กระดาษบนผนังย้ายเข้าระบบ — ESM เป็นที่ผลิตตัวเลข Actual ไม่ได้ทำแข่งกับ KPI Online ของกลุ่ม\n"
            "สถานะจริง: Action board ยัง 0 รายการ — ขอให้ผู้บริหารใช้ตั้ง action จริงอย่างน้อย 5 รายการ\n"
            "/group-overview เป็นตัวอย่างจอ (mockup) — ต้องบอกทุกครั้ง"),
]

# ═══════════════════════════ ตัววาด ═══════════════════════════
def build(lang):
    global LANG
    LANG = lang
    p = T.new_pres()
    page = [0]

    def notes(s, th):
        s.notes_slide.notes_text_frame.text = th

    def content(title, sub=None):
        s = T.blank(p); page[0] += 1
        T.head(s, t(title), t(sub) if sub else None)
        T.footer(s, page[0])
        return s

    def card(s, x, y, w, h, head, bullets, head_fill=GREEN_D, sz=12, bullet="•  "):
        T.rect(s, x, y, w, h, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
        T.rect(s, x, y, w, 0.44, fill=head_fill)
        T.text(s, x, y, w, 0.44, [[(t(head), {"sz": 14, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        T.text(s, x + 0.15, y + 0.56, w - 0.3, h - 0.64, [[(bullet + t(b), {"sz": sz, "c": GREEN})] for b in bullets], sa=5)

    def stat(s, x, y, w, big, label, sub=None):
        T.rect(s, x, y, w, 1.6, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
        T.text(s, x, y + 0.08, w, 0.7, [[(big, {"sz": 28 if len(big) <= 10 else 22, "b": True, "c": ORANGE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        T.text(s, x + 0.1, y + 0.8, w - 0.2, 0.42, [[(t(label), {"sz": 12, "b": True, "c": GREEN})]], align=PP_ALIGN.CENTER)
        if sub:
            T.text(s, x + 0.1, y + 1.2, w - 0.2, 0.35, [[(t(sub), {"sz": 10.5, "c": GREY})]], align=PP_ALIGN.CENTER)

    def table(s, x, y, colw, header, rows, sizes, bold_cols=(0,), color_cols=None, right_cols=()):
        tb = s.shapes.add_table(len(rows) + 1, len(colw), Inches(x), Inches(y), Inches(sum(colw)), Inches(0.38 * (len(rows) + 1))).table
        for c, w in enumerate(colw):
            tb.columns[c].width = Inches(w)
        for c, h in enumerate(header):
            cell = tb.cell(0, c); cell.fill.solid(); cell.fill.fore_color.rgb = GREEN_D; cell.text = t(h)
            r = cell.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(12.5); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = T.FONT
        for i, row in enumerate(rows, start=1):
            for c, v in enumerate(row):
                cell = tb.cell(i, c); cell.fill.solid(); cell.fill.fore_color.rgb = GREEN_TINT if i % 2 == 0 else WHITE
                cell.text = t(v); cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
                para = cell.text_frame.paragraphs[0]
                if c in right_cols: para.alignment = PP_ALIGN.RIGHT
                r = para.runs[0]; r.font.size = Pt(sizes[c]); r.font.name = T.FONT; r.font.bold = c in bold_cols
                r.font.color.rgb = (color_cols or {}).get(c, GREEN)

    def pillar_strip(s, items, y=5.45):
        w, gap = 2.98, 0.13
        for i, ((name, col), txt) in enumerate(zip(PILLARS, items)):
            x = 0.5 + i * (w + gap)
            T.rect(s, x, y, w, 1.02, fill=WHITE, line=col, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            T.rect(s, x, y, 0.12, 1.02, fill=col)
            T.text(s, x + 0.22, y + 0.06, w - 0.3, 0.3, [[(t(name), {"sz": 11, "b": True, "c": col})]])
            T.text(s, x + 0.22, y + 0.34, w - 0.3, 0.66, [[(t(txt), {"sz": 11, "c": GREEN})]])

    def flow_row(s, y, label, steps, color):
        T.text(s, 0.5, y - 0.45, 12, 0.4, [[(t(label), {"sz": 15, "b": True, "c": color})]])
        n = len(steps); w = (12.33 - 0.05 * (n - 1)) / n
        for i, (head, sub) in enumerate(steps):
            x = 0.5 + i * (w + 0.05)
            T.rect(s, x, y, w + 0.18, 0.95, fill=color if i % 2 == 0 else GREEN_D, shape=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON)
            T.text(s, x + 0.3, y + 0.05, w - 0.35, 0.85, [[(t(head), {"sz": 11.5, "b": True, "c": WHITE})], [(t(sub), {"sz": 9.5, "c": WHITE})]],
                   anchor=MSO_ANCHOR.MIDDLE, sa=1)

    # ── 1 ปก (R01: พื้นขาว) ──
    s = T.title_slide(p, t(COVER["title"]), t(COVER["sub"]), t(COVER["who"]), t(COVER["org"]),
                      decl_label=t(("My Quality Declaration", "คำประกาศคุณภาพของฉัน")))
    page[0] += 1
    notes(s, "เปิด: ESM คือระบบเดียวที่ทุกแผนกในโรงงานใช้ร่วมกัน — บันทึกครั้งเดียวที่หน้างาน แล้วทุกคนเห็นข้อมูลชุดเดียวกันทันที\n"
             "วันนี้จะเล่า 3 เรื่อง: แต่ละแผนกใช้อะไรได้บ้าง · งานของแต่ละแผนกเชื่อมกันยังไง · และได้อะไร 4 ข้อ: เห็นสด ลดงาน ใช้ง่าย ลดต้นทุน")

    # ── 2 Agenda (R01: ลิสต์เลขเขียว) ──
    page[0] += 1
    s = T.agenda_slide(p, page[0], t(COVER["title"]), t(("Agenda", "หัวข้อนำเสนอ")), [(t(a), t(b)) for a, b in AGENDA])
    notes(s, "5 ช่วง ใช้เวลารวม ~15 นาที: ภาพรวม → แผนกไหนใช้อะไร → งานเชื่อมกันยังไง → ลงรายละเอียดทีละแผนก → สรุปประโยชน์ 4 ข้อ")

    # ── 3 At a glance ──
    s = content(("ESM AT A GLANCE", "ESM ในหน้าเดียว"),
                ("Record once at the line — every department sees the same data, live", "บันทึกครั้งเดียวที่หน้างาน — ทุกแผนกเห็นข้อมูลชุดเดียวกันทันที"))
    for i, ((name, col), bl) in enumerate(zip(PILLARS, GLANCE if LANG == 0 else GLANCE_TH)):
        x = 0.5 + i * 3.11
        T.rect(s, x, 1.75, 2.98, 3.45, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
        T.rect(s, x, 1.75, 2.98, 0.62, fill=col)
        T.text(s, x, 1.75, 2.98, 0.62, [[(t(name), {"sz": 17, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        T.text(s, x + 0.18, 2.5, 2.65, 2.6, [[("✓  " + b, {"sz": 12.5, "c": GREEN})] for b in bl], sa=9)
    for i, (big, lab) in enumerate(GLANCE_STATS):
        x = 0.5 + i * 3.11
        T.rect(s, x, 5.38, 2.98, 0.9, fill=GREEN_TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        T.text(s, x + 0.1, 5.38, 0.95, 0.9, [[(big, {"sz": 24, "b": True, "c": ORANGE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        T.text(s, x + 1.05, 5.38, 1.85, 0.9, [[(t(lab), {"sz": 12, "b": True, "c": GREEN})]], anchor=MSO_ANCHOR.MIDDLE)
    T.text(s, 0.5, 6.4, 12.3, 0.35, [[(t(("Active users: ", "ผู้ใช้งานจริง: ")) + t(MEASURED), {"sz": 10, "i": True, "c": GREY})]])
    notes(s, "4 เสาหลักของ ESM — สีเดียวกันนี้จะโผล่ท้ายสไลด์ทุกแผนก (ส้ม=เห็นสด · เขียว=ลดงาน · เขียวเข้ม=ใช้ง่าย · เหลืองอำพัน=ลดต้นทุน)\n"
             "ตัวเลขผู้ใช้: บัญชี 100 · ใช้จริงใน 30 วัน 71 คน · ใน 7 วัน 47 คน (วัด 06/10)\n"
             "~1,300 ใบ/เดือน = ใบซ่อม MO 496 + ใบตรวจ PM 88 + รายงานกะ 505 + LPA 222 (ก.ย.)")

    # ── 4 Who uses what ──
    s = content(("WHO USES WHAT", "แผนกไหนใช้อะไร"),
                ("Every department works in the same system — each sees its own screens", "ทุกแผนกทำงานในระบบเดียวกัน แต่ละคนเห็นเฉพาะจอของตัวเอง"))
    table(s, 0.5, 1.7, [2.6, 9.73], [("Department", "แผนก"), ("Main screens / functions", "จอหลัก / ฟังก์ชัน")], MATRIX, [12, 11.5])
    notes(s, "ภาพเดียวตอบคำถาม \"แผนกไหนใช้อะไร\" — ทุกแผนกเข้าระบบเดียวกัน ด้วยบัญชีเดียว เมนูที่เห็นขึ้นกับสิทธิ์ของแต่ละคน\n"
             "ถ้าผู้บริหารถามว่าใช้จริงแค่ไหน: MTN/PD3/PD4 ใช้ทุกวัน · PD1/PD2/Store/Logistic/Planning กำลังขยาย · QA/JIG/DIE/PE ระบบพร้อม แต่ยังลงข้อมูลน้อย")

    # ── 5 How work connects ──
    s = content(("HOW THE WORK CONNECTS", "งานเชื่อมกันอย่างไร"),
                ("One event at the line moves through every department", "เหตุการณ์เดียวที่หน้างาน วิ่งผ่านทุกแผนกเอง"))
    flow_row(s, 2.15, FLOW1[0], FLOW1[1], ORANGE)
    flow_row(s, 3.9, FLOW2[0], FLOW2[1], GREEN)
    T.rect(s, 0.5, 5.25, 12.33, 1.0, fill=GREEN_TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, 0.7, 5.25, 12.0, 1.0, [[(t(TRACE[0]), {"sz": 13, "b": True, "c": GREEN_D}), (t(TRACE[1]), {"sz": 12.5, "c": GREEN})]], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "2 สายงานที่เห็นภาพชัดที่สุดว่าทุกแผนกต่อกันในระบบเดียว\n"
             "① เครื่องหยุด: ผลิตลง downtime → กดเรียกช่าง เลือกทีม (เด้งเฉพาะทีมนั้น · ไม่รับใน 15 นาทีเตือนเอง) → เปิดใบ MO ต่อจาก downtime → ซ่อม → QA ตรวจ → OEE กะนั้นอัปเดต → ขึ้น KPI บน OBEYA\n"
             "② ลูกค้าสั่งของ: เมล EDI เข้าคิวเอง → วางแผน/Monitoring → ไลน์สแกนคัมบังผลิต → ปิดใบแล้วระบบระเบิด BOM ขอพาร์ทลูกเอง → คลังนับแล้วรับ → ส่งของตัดสต็อกทันที\n"
             "ปิดด้วย /order-trace: เลขใบเดียวกางได้ทั้งเส้น")

    # ── 6 divider ──
    s = T.section_divider(p, "photo_press2500t.png", [t(("Department", "ลงรายละเอียด")), t(("by department", "ทีละแผนก"))]); page[0] += 1
    notes(s, "ต่อไปลงรายละเอียดทีละแผนก — ทุกสไลด์มี 3 ช่องเหมือนกัน: ทำอะไรได้ · จอหลัก · เชื่อมกับใคร และแถบ 4 สีด้านล่าง")

    # ── 7–16 แผนก ──
    for d in DEPTS:
        s = content(d["title"], d["sub"])
        y, h = 1.72, 3.58
        card(s, 0.5, y, 5.35, h, ("WHAT YOU CAN DO", "ทำอะไรได้บ้าง"), d["feat"], sz=12.5)
        card(s, 6.0, y, 3.25, h, ("KEY SCREENS", "จอหลัก"), d["scr"], head_fill=GREEN, sz=12.5, bullet="•  ")
        x, w = 9.4, 3.43
        T.rect(s, x, y, w, h, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
        T.rect(s, x, y, w, 0.44, fill=ORANGE)
        T.text(s, x, y, w, 0.44, [[(t(("CONNECTED TO", "เชื่อมกับแผนก")), {"sz": 14, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        for i, (who, what) in enumerate(d["links"]):
            yy = y + 0.6 + i * 0.73
            T.rect(s, x + 0.15, yy, 1.05, 0.34, fill=GREEN_TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            T.text(s, x + 0.15, yy, 1.05, 0.34, [[(t(who), {"sz": 10.5, "b": True, "c": GREEN_D})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            T.text(s, x + 1.27, yy - 0.04, w - 1.4, 0.7, [[(t(what), {"sz": 10.5, "c": GREEN})]])
        pillar_strip(s, d["pil"])
        notes(s, d["th"])

    # ── divider ──
    s = T.section_divider(p, "photo_robots.png", [t(("Why it", "ได้ประโยชน์")), t(("matters", "อะไรบ้าง"))]); page[0] += 1
    notes(s, "สรุปประโยชน์ 4 ข้อ พร้อมตัวเลขที่วัดได้จริง")

    # ── Realtime ──
    s = content(("REAL-TIME: EVERYONE SEES THE SAME MOMENT", "เห็นทันที: ทุกคนเห็นเหตุการณ์เดียวกันพร้อมกัน"),
                ("From 'we find out at month-end' to 'we see it during the shift'", "จาก 'รู้ตอนสิ้นเดือน' เป็น 'เห็นระหว่างกะ'"))
    for i, (big, lab, sub) in enumerate([
            ("529 → 30 min", ("Call → acknowledged", "เรียกช่าง → รับเรื่อง"), ("P80 · Aug → Sep", "P80 · ส.ค. → ก.ย.")),
            ("7.7 → 3.2 min", ("Stop → repair order", "เครื่องหยุด → เปิดใบซ่อม"), ("median · Aug → Sep", "ค่ากลาง · ส.ค. → ก.ย.")),
            ("7.9 → 0.9 h", ("Repair order → done", "เปิดใบซ่อม → ซ่อมเสร็จ"), ("P80 · Aug → Sep", "P80 · ส.ค. → ก.ย.")),
            ("70% → 96%", ("Shifts with OEE at close", "กะที่มี OEE ตอนปิดกะ"), ("Jul → Sep", "ก.ค. → ก.ย."))]):
        if LANG == 1:
            big = big.replace(" min", " นาที").replace(" h", " ชม.")
        stat(s, 0.5 + i * 3.12, 1.75, 2.95, big, lab, sub)
    card(s, 0.5, 3.55, 6.1, 2.75, ("WHAT UPDATES LIVE", "อะไรอัปเดตสดบ้าง"), [
        ("Andon colour of every line — factory map and TV", "สถานะ Andon ทุกไลน์ บนผังโรงงานและจอ TV"),
        ("OEE, output and stops during the shift", "OEE ยอดผลิต และเครื่องหยุดระหว่างกะ"),
        ("Maintenance Andon with siren per team", "บอร์ด Andon ช่างพร้อมเสียงเตือนแยกทีม"),
        ("Shipping rounds: enough / short / late", "รอบส่งของ: พอ / ขาด / ล่าช้า")], head_fill=ORANGE, sz=12.5)
    card(s, 6.73, 3.55, 6.1, 2.75, ("HOW PEOPLE ARE TOLD", "แจ้งคนที่เกี่ยวข้องอย่างไร"), [
        ("Phone push and Telegram to the right team only", "แจ้งเตือนมือถือและ Telegram เฉพาะทีมที่เกี่ยวข้อง"),
        ("Escalates by itself when nobody answers in 15 min", "ไม่มีคนรับใน 15 นาที ระบบเตือนต่อเอง"),
        ("'My queue' — what waits for me, in one place", "'คิวงานของฉัน' งานที่รอเราอยู่ที่เดียว"),
        ("Morning summary per department", "สรุปงานตอนเช้ารายแผนก")], head_fill=ORANGE, sz=12.5)
    T.text(s, 0.5, 6.4, 12.3, 0.35, [[(t(MEASURED), {"sz": 10, "i": True, "c": GREY})]])
    notes(s, "Realtime: ตัวเลข 4 ตัวบนสุดวัดจากฐานข้อมูลจริง 06/10 (เรียกช่าง 279 ครั้ง · MO 294/456 ใบ · กะ 485 กะ)\n"
             "หลักคิด: ระบบไม่ได้ยิงแจ้งเตือนทุกคน — เด้งเฉพาะคนที่ต้องทำ ไม่งั้นคนจะเลิกดู")

    # ── Less work ──
    s = content(("LESS WORK: PAPER AND RE-TYPING REMOVED", "ลดภาระงาน: เลิกกระดาษและการคีย์ซ้ำ"),
                ("Records per month that used to be paper or Excel", "จำนวนบันทึกต่อเดือนที่เมื่อก่อนเป็นกระดาษหรือ Excel"))
    recs = [(("Repair orders (MO)", "ใบแจ้งซ่อม (MO)"), "FM-MTN-006 / FM-JIG-008", "76", "496"),
            (("PM inspections", "ใบตรวจ PM"), ("Paper check sheets", "Check sheet กระดาษ"), "6", "88"),
            (("Downtime records", "บันทึกเครื่องหยุด"), ("Notebook / whiteboard", "สมุด / กระดาน"), "3,688", "3,792"),
            (("Shift reports", "รายงานกะ"), ("Paper daily report", "Daily report กระดาษ"), "497", "505"),
            (("LPA audits", "ตรวจ LPA"), "FM-QMR-008", "289", "222"),
            (("Check-in records", "บันทึกเช็คชื่อ"), ("Attendance sheet", "ใบเช็คชื่อ"), "2,897", "3,489")]
    table(s, 0.5, 1.75, [2.4, 2.6, 1.15, 1.15], [("Record", "บันทึก"), ("Replaces", "แทนของเดิม"), ("Aug", "ส.ค."), ("Sep", "ก.ย.")], recs,
          [12, 12, 12, 12], bold_cols=(0, 3), color_cols={3: ORANGE}, right_cols=(2, 3))
    card(s, 8.0, 1.75, 4.83, 3.9, ("THE SYSTEM DOES IT", "ระบบทำให้เอง"), [
        ("Calculates OEE, MTTR, MTBF, PPM", "คำนวณ OEE · MTTR · MTBF · PPM"),
        ("Works out PM due dates and kanban", "คิดวันครบกำหนด PM และจำนวนคัมบัง"),
        ("Opens follow-up orders (MO, parts, purchase)", "เปิดใบงานต่อเนื่อง (MO · ขอพาร์ท · ขอซื้อ)"),
        ("Prints official FM- forms with running numbers", "พิมพ์ฟอร์ม FM- ทางการพร้อมเลขที่เอกสาร"),
        ("Exports monthly KPI packs", "ส่งออกชุด KPI รายเดือน")], sz=12.5)
    T.rect(s, 0.5, 5.8, 12.33, 0.55, fill=GREEN_TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, 0.7, 5.8, 12.0, 0.55, [[(t(("Record once at the line → no re-typing into Excel, no copying between departments.",
                                          "บันทึกครั้งเดียวที่หน้างาน → ไม่ต้องคีย์ลง Excel ซ้ำ ไม่ต้องส่งต่อข้อมูลกันเองระหว่างแผนก")),
                                        {"sz": 13, "b": True, "c": GREEN_D})]], anchor=MSO_ANCHOR.MIDDLE)
    T.text(s, 0.5, 6.4, 12.3, 0.35, [[(t(MEASURED), {"sz": 10, "i": True, "c": GREY})]])
    notes(s, "ลดงาน: ตารางคือจำนวนบันทึกต่อเดือนที่เมื่อก่อนเป็นกระดาษ/Excel (วัด 06/10)\n"
             "MO 76 → 496 และ PM 6 → 88 คือแผนกที่ย้ายจากกระดาษมาเต็มตัวในเดือนเดียว\n"
             "LPA ก.ย. ลดลง (289 → 222) — ถ้าถูกถาม: เป็นจำนวนการตรวจ ไม่ใช่ปัญหาระบบ")

    # ── Easy to use ──
    s = content(("EASY TO USE: BUILT FOR THE SHOPFLOOR", "ใช้งานง่าย: ออกแบบมาเพื่อหน้างาน"),
                ("Designed with the people who use it every day", "ออกแบบร่วมกับคนที่ใช้งานทุกวัน"))
    for i, (h, d) in enumerate([
            (("Scan, don't search", "สแกน ไม่ต้องค้นหา"), ("QR labels on machines, jigs, parts and delivery points — scan to pick",
                                                         "ป้าย QR ติดเครื่องจักร จิ๊ก ชิ้นงาน และจุดส่ง สแกนแล้วเลือกให้เลย")),
            (("Thai screens, one login", "จอภาษาไทย บัญชีเดียว"), ("Same account for every department; menu shows only what you need",
                                                                  "บัญชีเดียวทุกแผนก เมนูโชว์เฉพาะที่ต้องใช้")),
            (("Phone, PC or factory TV", "มือถือ PC หรือจอ TV"), ("Tested on phone width and on the 2023 TV browsers at the lines",
                                                                "ทดสอบบนจอมือถือ และเบราว์เซอร์ของจอ TV รุ่นปี 2023 ที่หน้าไลน์")),
            (("Control the TV from a phone", "ใช้มือถือคุมจอ TV"), ("No keyboard needed at the board — pair with a 6-digit code",
                                                                 "ไม่ต้องมีคีย์บอร์ดที่จอ จับคู่ด้วยรหัส 6 หลัก")),
            (("'My queue'", "'คิวงานของฉัน'"), ("Everything waiting for me in one list — nothing to search for",
                                               "งานที่รอเราทั้งหมดในลิสต์เดียว ไม่ต้องไล่หา")),
            (("Mistakes caught early", "กันพลาดตั้งแต่ต้น"), ("Wrong part blocked · time outside the shift warned · AM/PM swap fixed in one tap",
                                                            "หยิบผิดชิ้นถูกบล็อก · เวลาหลุดกะถูกเตือน · AM/PM สลับแก้ได้คลิกเดียว"))]):
        r_, c_ = divmod(i, 3)
        x, yy = 0.5 + c_ * 4.15, 1.75 + r_ * 2.3
        T.rect(s, x, yy, 4.0, 2.12, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
        T.rect(s, x, yy, 4.0, 0.12, fill=GREEN_D)
        T.text(s, x + 0.2, yy + 0.28, 3.6, 0.5, [[(t(h), {"sz": 16, "b": True, "c": GREEN_D})]])
        T.text(s, x + 0.2, yy + 0.85, 3.6, 1.2, [[(t(d), {"sz": 12.5, "c": GREEN})]])
    notes(s, "ใช้ง่าย: ออกแบบจาก feedback หน้างาน เช่น สแกน QR แทนเลือกจาก dropdown ยาวๆ ตอนใส่ถุงมือ\n"
             "จอ TV หน้างานเป็นเบราว์เซอร์รุ่นเก่า (Chromium 94 บน webOS 2023) — ระบบทดสอบให้ผ่านทุกครั้งก่อนขึ้นใช้งาน")

    # ── Lower cost ──
    s = content(("LOWER COST: WHERE THE SAVINGS COME FROM", "ลดต้นทุน: ประหยัดได้จากตรงไหน"),
                ("Operational effects measured today — baht value to be calculated with Finance", "ผลที่วัดได้วันนี้ — มูลค่าเป็นเงินบาทจะคำนวณร่วมกับฝ่ายบัญชี"))
    for i, (h, d, big) in enumerate([
            (("No licence fees", "ไม่มีค่าไลเซนส์"), ("Built in-house by the plant team — no per-user or per-module licence", "ทีมในโรงงานพัฒนาเอง ไม่มีค่าไลเซนส์รายคนหรือรายโมดูล"), ("In-house", "พัฒนาเอง")),
            (("Hardware we already own", "ใช้อุปกรณ์ที่มีอยู่"), ("Runs in a web browser on existing PCs, phones and line TVs", "ทำงานบนเว็บเบราว์เซอร์ ใช้ PC มือถือ และจอ TV ที่มีอยู่"), ("Existing devices", "อุปกรณ์เดิม")),
            (("Less paper & printing", "ลดกระดาษและงานพิมพ์"), ("Repair, PM, shift and LPA forms now digital", "ใบซ่อม ใบตรวจ PM รายงานกะ และ LPA อยู่ในระบบแล้ว"), ("~1,300 forms / month", "~1,300 ใบ/เดือน")),
            (("Less machine waiting", "เครื่องรอซ่อมน้อยลง"), ("Repair order → done, P80", "เปิดใบซ่อม → ซ่อมเสร็จ (P80)"), ("7.9 h → 0.9 h", "7.9 → 0.9 ชม.")),
            (("Less admin time", "ลดเวลางานเอกสาร"), ("No re-typing paper into Excel; KPI and forms generated", "ไม่ต้องคีย์กระดาษลง Excel KPI และฟอร์มสร้างให้เอง"), ("Record once", "บันทึกครั้งเดียว")),
            (("Low-cost cloud", "คลาวด์ต้นทุนต่ำ"), ("No server room; usage-based cloud with image and data limits", "ไม่ต้องมีห้องเซิร์ฟเวอร์ จ่ายตามการใช้งาน คุมขนาดรูปและข้อมูล"), ("Pay for use", "จ่ายตามใช้"))]):
        r_, c_ = divmod(i, 3)
        x, yy = 0.5 + c_ * 4.15, 1.75 + r_ * 2.2
        T.rect(s, x, yy, 4.0, 2.02, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
        T.rect(s, x, yy, 0.12, 2.02, fill=AMBER)
        bt = t(big)
        T.text(s, x + 0.3, yy + 0.15, 3.55, 0.5, [[(bt, {"sz": 19 if len(bt) <= 14 else 16, "b": True, "c": ORANGE})]])
        T.text(s, x + 0.3, yy + 0.72, 3.55, 0.4, [[(t(h), {"sz": 14, "b": True, "c": GREEN_D})]])
        T.text(s, x + 0.3, yy + 1.15, 3.55, 0.85, [[(t(d), {"sz": 11.5, "c": GREEN})]])
    T.text(s, 0.5, 6.25, 12.3, 0.5, [[(t(MEASURED) + t((" · ~1,300 = MO 496 + PM 88 + shift reports 505 + LPA 222 (Sep)",
                                                       " · ~1,300 = MO 496 + PM 88 + รายงานกะ 505 + LPA 222 (ก.ย.)")), {"sz": 10, "i": True, "c": GREY})]])
    notes(s, "ลดต้นทุน — สำคัญ: สไลด์นี้ไม่ใส่ตัวเลขเงินบาท เพราะยังไม่ได้คำนวณร่วมกับฝ่ายบัญชี (ห้ามเดาตัวเลขต่อหน้าผู้บริหาร)\n"
             "พูดได้: ไม่มีค่าไลเซนส์ (ทำเองในโรงงาน) · ใช้อุปกรณ์เดิม · กระดาษ ~1,300 ใบ/เดือนย้ายเข้าระบบ · เครื่องรอซ่อมสั้นลงชัดเจน\n"
             "ถ้าถูกถามตัวเลขเงิน: เสนอว่าจะคำนวณร่วมกับบัญชี เช่น ชั่วโมงเครื่องรอซ่อมที่ลดลง × ต้นทุนต่อชั่วโมงของไลน์")

    # ── Next steps ──
    s = content(("NEXT STEPS", "ก้าวต่อไป"),
                ("Deeper use at TSAT first — then a recipe for other plants", "ใช้ให้ลึกที่ TSAT ก่อน แล้วทำสูตรขยายไปโรงงานอื่น"))
    for i, (h, d) in enumerate([
            (("Every department records daily", "ทุกแผนกบันทึกทุกวัน"), ("QA bins, JIG / DIE status, PE documents and planning — from 'ready' to 'used'",
                                                                       "ถัง QA · สถานะ JIG/DIE · เอกสาร PE · แผนผลิต — จาก 'พร้อมใช้' เป็น 'ใช้จริง'")),
            (("Managers act on the numbers", "ผู้บริหารใช้ตัวเลขตัดสินใจ"), ("Action board in OBEYA: at least 5 real actions per month with owner and date",
                                                                           "ตั้ง Action ใน OBEYA อย่างน้อยเดือนละ 5 รายการ มีผู้รับผิดชอบและกำหนดเสร็จ")),
            (("Clear the slow hand-offs", "แก้จุดส่งต่อที่ยังช้า"), ("Repair orders waiting for QA · stock receipts after closing orders",
                                                                   "ใบซ่อมรอ QA ตรวจ · การรับเข้าคลังหลังปิดใบผลิต")),
            (("Connect machines (later)", "ต่อเครื่องจักร (เฟสถัดไป)"), ("SCADA signals as facts; people still give the reasons and the NG",
                                                                       "สัญญาณ SCADA เป็นข้อเท็จจริง คนยังเป็นผู้บอกสาเหตุและของเสีย")),
            (("Roll out to other plants", "ขยายไปโรงงานอื่น"), ("Same system, own data — start with Production + MTN, then add departments",
                                                              "ระบบเดียวกัน ข้อมูลแยกแต่ละโรง เริ่มที่ผลิตและซ่อมบำรุง แล้วค่อยเพิ่มแผนก"))]):
        yy = 1.7 + i * 0.95
        T.rect(s, 0.5, yy, 0.62, 0.62, fill=ORANGE, shape=MSO_SHAPE.OVAL)
        T.text(s, 0.5, yy, 0.62, 0.62, [[(str(i + 1), {"sz": 18, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        T.rect(s, 1.3, yy - 0.06, 11.53, 0.78, fill=GREEN_TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        T.text(s, 1.5, yy - 0.06, 3.8, 0.78, [[(t(h), {"sz": 14, "b": True, "c": GREEN_D})]], anchor=MSO_ANCHOR.MIDDLE)
        T.text(s, 5.3, yy - 0.06, 7.4, 0.78, [[(t(d), {"sz": 12.5, "c": GREEN})]], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "ก้าวต่อไป: (1) ทุกแผนกลงข้อมูลทุกวัน (2) ผู้บริหารใช้ตัวเลขตั้ง action จริง (3) แก้จุดส่งต่อที่ยังช้า — MO รอ QA 205 ใบ, รับเข้าคลัง ~73 ชม.\n"
             "(4) ต่อเครื่องจักร SCADA ภายหลัง — เครื่องบอกข้อเท็จจริง คนบอกเหตุผล (5) ขยายไปโรงงานอื่น เริ่มที่ผลิต + ซ่อมบำรุงก่อน\n"
             "สิ่งที่ยังไม่ได้ทำ ห้ามพูดว่าเสร็จ: ต่อ SAP/SCADA · ระบบหลายบริษัทจริง (/group-overview เป็นตัวอย่างจอ)")

    # ── ปิด (R01: พื้นขาว) ──
    page[0] += 1
    s = T.closing_slide(p, page[0], "Before We Build Parts, We Build People", [
        (t(("Record once. Everyone sees it. Live.", "บันทึกครั้งเดียว ทุกคนเห็น ทันที")), ORANGE),
        (t(("Thank you  ·  Q & A", "ขอบคุณ  ·  ถาม-ตอบ")), GREEN)])
    notes(s, "ปิด: บันทึกครั้งเดียว ทุกคนเห็น และเห็นทันที — ระบบไม่ได้ทำงานแทนคน แต่ทำให้ทุกแผนกส่งงานต่อกันเร็วขึ้นบนข้อมูลชุดเดียวกัน")

    out = os.path.join(HERE, "..", f"TSAT_ESM_Program_Overview_2026-10_{'EN' if lang == 0 else 'TH'}.pptx")
    p.save(out)
    print("saved", os.path.abspath(out), "slides:", len(p.slides))


if __name__ == "__main__":
    build(0)
    build(1)

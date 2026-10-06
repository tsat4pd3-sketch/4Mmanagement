# -*- coding: utf-8 -*-
"""ESM Department Roadshow deck (TSG theme) — draft 1 · 2026-10-06
Run:  python3 build_roadshow.py   →  ../TSAT_ESM_Department_Roadshow_2026-10.pptx
Sources (every number traced):
  · roadshow-metrics-2026-10-06.md  (DB queries 06/10 — Production→MTN · shift close)
  · roadshow-dept-stories-2026-10-06.md (before/after per department, cited to docs)
  · ADOPTION-ROADMAP.md (status 05/10)
Rule: slide text = English (TSG checklist) · speaker notes = Thai script for presenters.
"Before ESM" numbers do NOT exist in the system ⇒ shown as FILL-IN boxes for each department.
"""
import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import tsg_theme as T
from tsg_theme import GREEN_D, GREEN_M, GREEN_TINT, ORANGE, GREY, AMBER, WHITE, GOLD, RGBColor

PART = RGBColor(0xD8, 0xE4, 0xD0)
MEASURED = "Measured from ESM database · 06 Oct 2026"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "TSAT_ESM_Department_Roadshow_2026-10.pptx")

p = T.new_pres()
page = [1]


def notes(s, th):
    s.notes_slide.notes_text_frame.text = th


def content(title, sub=None):
    s = T.blank(p)
    T.head(s, title, sub)
    page[0] += 1
    T.footer(s, page[0])
    return s


def card(s, x, y, w, h, head, bullets, head_fill=GREEN_D, sz=12.5):
    T.rect(s, x, y, w, h, fill=WHITE, line=PART, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    T.rect(s, x, y, w, 0.46, fill=head_fill)
    T.text(s, x, y, w, 0.46, [[(head, {"sz": 15, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    runs = [[("•  " + b, {"sz": sz, "c": GREEN_M})] for b in bullets]
    T.text(s, x + 0.15, y + 0.6, w - 0.3, h - 0.7, runs, sa=6)


def chip(s, x, y, label, color):
    T.rect(s, x, y, 2.3, 0.38, fill=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, x, y, 2.3, 0.38, [[(label, {"sz": 12, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def fill_in(s, x, y, w):
    T.rect(s, x, y, w, 0.5, fill=None, line=AMBER, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, x + 0.15, y, w - 0.3, 0.5, [[("Department to fill:  time per day BEFORE  ____ min   →   WITH ESM  ____ min   ·   forms per month removed  ____",
                                             {"sz": 12, "b": True, "c": AMBER})]], anchor=MSO_ANCHOR.MIDDLE)


def stat(s, x, y, w, big, label, sub=None):
    T.rect(s, x, y, w, 1.75, fill=WHITE, line=PART, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    T.text(s, x, y + 0.12, w, 0.75, [[(big, {"sz": 30 if len(big) <= 10 else 24, "b": True, "c": ORANGE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    T.text(s, x + 0.1, y + 0.9, w - 0.2, 0.45, [[(label, {"sz": 12.5, "b": True, "c": GREEN_D})]], align=PP_ALIGN.CENTER)
    if sub:
        T.text(s, x + 0.1, y + 1.3, w - 0.2, 0.4, [[(sub, {"sz": 11, "c": GREY})]], align=PP_ALIGN.CENTER)


def dept(title, sub, status, status_color, before, after, cut, handoff, th):
    s = content(title, sub)
    chip(s, 10.55, 1.17, status, status_color)
    y, h = 1.75, 3.7
    card(s, 0.5, y, 4.0, h, "BEFORE  (manual)", before, head_fill=GREY)
    card(s, 4.67, y, 4.0, h, "WITH ESM", after)
    card(s, 8.84, y, 4.0, h, "WHAT WE CUT", cut, head_fill=ORANGE)
    T.rect(s, 0.5, 5.58, 12.34, 0.62, fill=GREEN_TINT)
    T.text(s, 0.65, 5.58, 12.1, 0.62, [[("Works with others:  ", {"sz": 12.5, "b": True, "c": GREEN_D}), (handoff, {"sz": 12, "c": GREEN_M})]], anchor=MSO_ANCHOR.MIDDLE)
    fill_in(s, 0.5, 6.3, 12.34)
    notes(s, th)
    return s


# ── 1 Title ─────────────────────────────────────────────────────────
s = T.title_slide(p, "ESM Department Roadshow",
                  "From paper to a connected shopfloor",
                  "Department roadshow · each department presents its own floor",
                  "TSAT · Thai Summit Group · October 2026",
                  "Zero defect is possible")
notes(s, "เปิด: วันนี้ไม่ใช่การโชว์ฟีเจอร์ แต่แต่ละส่วนงานจะเล่าว่า เมื่อก่อนทำงานแบบ manual ยังไง พอใช้ ESM แล้วลดอะไร ช่วยอะไร และการทำงานกับแผนกอื่นไวขึ้นแค่ไหน — ตัวเลขทุกตัววัดจากฐานข้อมูลจริง วันที่เขียนไว้มุมสไลด์")

# ── 2 Agenda ────────────────────────────────────────────────────────
s = T.blank(p); page[0] += 1
T.text(s, 0.5, 0.28, 8, 0.95, [[("Agenda", {"sz": 34, "b": True, "c": GREEN_D})]], anchor=MSO_ANCHOR.MIDDLE)
s.shapes.add_picture(os.path.join(T.ASSET, "logo_big.png"), Inches(9.2), Inches(1.6), Inches(3.4), Inches(3.6))
items = [("Why ESM", "One shopfloor record instead of paper, Excel and chat groups"),
         ("Cross-department speed", "Measured hand-offs: calls, repairs, approvals"),
         ("Department stories", "Production · MTN · JIG/DIE · Store · Logistic · Planning · QA · PE · HR · Management"),
         ("Honest status & next steps", "What is still slow and what each department commits to")]
for i, (t1, t2) in enumerate(items):
    yy = 1.55 + i * 1.18
    T.rect(s, 0.7, yy, 0.62, 0.62, fill=ORANGE, shape=MSO_SHAPE.OVAL)
    T.text(s, 0.7, yy, 0.62, 0.62, [[(str(i + 1), {"sz": 20, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    T.text(s, 1.55, yy - 0.05, 7.4, 0.45, [[(t1, {"sz": 18, "b": True, "c": GREEN_D})]])
    T.text(s, 1.55, yy + 0.38, 7.4, 0.4, [[(t2, {"sz": 12, "c": GREY})]])
T.footer(s, page[0])
notes(s, "4 ช่วง: ทำไมต้อง ESM → ตัวเลขการส่งต่องานข้ามแผนกที่วัดได้จริง → แต่ละส่วนงานเล่าหน้างานตัวเอง (ส่วนงานละ ~3 นาที) → สถานะตรงๆ ว่าอะไรยังช้า และแต่ละส่วนงานจะทำอะไรต่อ")

# ── 3 Why ESM ───────────────────────────────────────────────────────
s = content("Why ESM", "Information used to live in paper, Excel and chat groups")
card(s, 0.5, 1.8, 6.0, 4.3, "BEFORE", [
    "Paper forms written by hand, then re-typed into Excel",
    "Calling maintenance = walking or phoning",
    "Requesting parts from Store = LINE group chat",
    "OEE known only at month-end, if at all",
    "Tracing a problem = searching files and asking shift leaders",
], head_fill=GREY, sz=14)
card(s, 6.84, 1.8, 6.0, 4.3, "WITH ESM", [
    "Record once at the line — every department sees the same data",
    "Calls, repair orders, approvals flow to the right team automatically",
    "Live OEE during the shift, locked when the shift is approved",
    "Printed forms (FM-…) generated from the data, numbered by the system",
    "One order number traces the whole chain",
], sz=14)
notes(s, "ประโยคหลัก: \"บันทึกครั้งเดียวที่หน้างาน แล้วทุกแผนกเห็นข้อมูลชุดเดียวกัน\" · ตัวอย่างจากหน้างานจริง: หัวหน้ากลุ่ม Assy2 เคยต้องเขียนมือ 3 ใบทุกครั้งทั้งที่ข้อมูลอยู่ในระบบแล้ว · สโตร์เคยอ่านจาก LINE ว่าใครเบิก (\"มั่วมาก\" — ทีมสโตร์เล่าเอง)")

# ── 4 Cross-department speed (measured) ─────────────────────────────
s = content("Cross-department hand-offs got faster", "First month live vs September 2026 — " + MEASURED)
stat(s, 0.5, 1.85, 2.95, "529 → 30 min", "Call technician → acknowledged", "P80 · Aug → Sep · 279 calls")
stat(s, 3.62, 1.85, 2.95, "7.7 → 3.2 min", "Machine stop → repair order opened", "median · Aug → Sep · 294 MOs")
stat(s, 6.74, 1.85, 2.95, "7.9 → 0.9 h", "Repair order → repair done", "P80 · Aug → Sep · 456 MOs")
stat(s, 9.86, 1.85, 2.95, "70% → 96%", "Shifts with OEE ready at close", "Jul → Sep · 485 shifts")
T.headline_box(s, 0.5, 3.9, 12.31, "What changed in the flow")
for i, t in enumerate([
    "Production presses 📞 and picks the team (JIG / DIE / MTN) — only that team is alerted; unanswered 15 min = escalation",
    "A repair order opens straight from the downtime record — machine, line and cost center are filled in, nothing re-typed",
    "Closing a repair closes the downtime — OEE for the shift is correct immediately",
]):
    T.text(s, 0.7, 4.45 + i * 0.55, 12.0, 0.5, [[("✔  " + t, {"sz": 13, "c": GREEN_M})]], anchor=MSO_ANCHOR.MIDDLE)
notes(s, "ตัวเลขทั้ง 4 วัดจากฐานข้อมูลจริง 06/10 เทียบเดือนแรกที่มีข้อมูลกับ ก.ย. 2026 (ไม่ใช่เทียบกับยุคกระดาษ เพราะยุคกระดาษไม่มีข้อมูล)\n"
          "• ช่างรับทราบการเรียก P80: 529 นาที (ส.ค.) → 30 นาที (ก.ย.) — 8 ใน 10 ครั้ง ช่างรับทราบภายใน 30 นาที\n"
          "• เครื่องหยุด → เปิดใบซ่อม ค่ากลาง 7.7 → 3.2 นาที\n"
          "• ใบซ่อม → ซ่อมเสร็จ P80: 7.9 ชม. → 0.9 ชม.\n"
          "• กะที่มี OEE พร้อมตอนปิดกะ: 70% → 96%\n"
          "ข้อควรพูดตรงๆ: ค่ากลางเวลารับทราบเดือน ต.ค. ขยับขึ้นเล็กน้อย (6.2 → 9.0 นาที) — ข้อมูล 6 วันแรกยังน้อย")

# ── 5 Hand-off map ──────────────────────────────────────────────────
s = content("Hand-off map — who works with whom", "Before (manual) vs with ESM · status as of 05 Oct 2026")
rows = [
    ("Production → MTN", "Walk / phone the technician", "📞 call by team · MO opened from downtime", "In use"),
    ("MTN → QA → requester", "Paper form passed by hand", "Step-by-step MO, morning digest 'waiting QA'", "Slow at QA"),
    ("Production → Store", "LINE group chat", "WIP call → pick scan → delivery scan → receive", "Starting"),
    ("Close FG order → Store / Purchasing", "Calculated by hand", "BOM explodes automatically into lot / buy / raw requests", "Auto works"),
    ("Customer → Sales", "Save mail attachments, upload by hand", "EDI 830/862 mail queued automatically", "In use"),
    ("Planning → Production", "13-sheet Excel updated daily", "/monitoring + /production-plan", "Starting"),
    ("Production → QA (bins)", "3 hand-written forms", "Bin from the shift record, QA decides in 4 ways", "Not yet"),
    ("MTN → Production (PM)", "Plan sent by e-mail", "PM forecast + coordination board", "In use"),
    ("Meeting → everyone", "In heads / chat / paper", "Action board with owner and due date", "Not yet"),
]
x0, y0, cw = 0.5, 1.75, [2.9, 2.9, 4.9, 1.6]
hdr = ["Hand-off", "Before (manual)", "With ESM", "Status"]
xx = x0
for j, t in enumerate(hdr):
    T.rect(s, xx, y0, cw[j], 0.4, fill=GREEN_D)
    T.text(s, xx + 0.08, y0, cw[j] - 0.1, 0.4, [[(t, {"sz": 12, "b": True, "c": WHITE})]], anchor=MSO_ANCHOR.MIDDLE)
    xx += cw[j]
for i, r in enumerate(rows):
    yy = y0 + 0.4 + i * 0.5
    xx = x0
    for j, t in enumerate(r):
        T.rect(s, xx, yy, cw[j], 0.5, fill=GREEN_TINT if i % 2 == 0 else WHITE)
        col = GREEN_M
        if j == 3:
            col = GREEN_M if t in ("In use", "Auto works") else (AMBER if t in ("Starting", "Slow at QA") else ORANGE)
        T.text(s, xx + 0.08, yy, cw[j] - 0.12, 0.5, [[(t, {"sz": 11.5, "b": j in (0, 3), "c": col})]], anchor=MSO_ANCHOR.MIDDLE)
        xx += cw[j]
notes(s, "ตารางนี้คือคำตอบของคำถาม \"workflow ที่ต้องเกี่ยวข้องกับคนอื่นดีขึ้นยังไง\" — คอลัมน์สถานะพูดตรงๆ: ขาที่ระบบทำเองทำงานครบ แต่ขาที่ต้องมีคนกด (QA ตัดสิน, สโตร์รับใบ, ประชุมออกใบติดตาม) ยังค้าง")

# ── Department slides ───────────────────────────────────────────────
T_LIVE, T_PART, T_NOT = ("Live", GREEN_M), ("Partly live", AMBER), ("Not started", ORANGE)

dept("Production (PD1–PD4)", "From hand-written forms to live OEE", *T_LIVE,
     ["Group leader hand-writes problem report, bin tag and scrap report (3 forms)",
      "Morning meeting needs someone to make slides",
      "LPA plan and BBS on paper / Excel",
      "OEE not known until month-end"],
     ["Scan kanban to open/close orders; RH/LH pairs open together",
      "Live OEE during the shift, locked after approval",
      "Problem report & scrap report (FM-PD2-002) built from the shift data",
      "Morning meeting agenda generated automatically"],
     ["Re-typing the same data into 3 forms",
      "Meeting slide preparation",
      "Daily LPA planning by hand (fills a whole quarter)",
      "Arguing which OEE number is right"],
     "calls MTN by team · requests parts from Store · sends bins to QA · receives PM plan from MTN",
     "ผลิต: ยกตัวอย่างหัวหน้ากลุ่ม Assy2 ที่เคยเขียนมือ 3 ใบ (FM-PD1-019 / FM-PD2-002 / FM-PD2-023) — ตอนนี้กดปุ่มเดียวจากข้อมูลกะ\n"
     "Demo: /daily-report ไลน์ 060/061 → OEE สด + %Q 100 ที่ยืนยันแล้ว → /morning-meeting โหมด TV แล้วกด ➕ Action สร้างใบติดตามจริง 1 ใบ\n"
     "ห้ามพูดเกิน: PD1/PD2 เพิ่งเริ่ม 17/09 ยังโชว์ OEE ไม่ได้ · งานแก้ (rework) ยังแทบไม่มีคนบันทึก · ปิดกะยังไม่บังคับตอบเรื่องคุณภาพ\n"
     "ให้หัวหน้าส่วนงานกรอกกล่องสีเหลือง: เมื่อก่อนใช้เวลาเขียนใบ/ทำสไลด์วันละกี่นาที")

dept("Maintenance (MTN)", "Repair orders that move by themselves", *T_LIVE,
     ["Paper repair form FM-MTN-006 passed by hand",
      "Spare-part list in Excel, rank decided by hand",
      "Monthly summary = add up printed forms",
      "Every call reached all 13 technicians"],
     ["Scan machine QR → line, section, cost center filled in",
      "Repair can start at once; manager signs after (line does not wait)",
      "Response / repair-time clocks start automatically",
      "Spare parts deducted in the MO, A/B/C rank automatic"],
     ["Waiting for approval before repairing",
      "Calls to the wrong team",
      "Building the monthly report by hand (Excel export, 3 sheets)",
      "Ranking spare parts by hand"],
     "receives calls from Production · hands over to QA and the requester · sends PM plan to Production",
     "MTN: ตัวเลขจริง ก.ย. — ช่างรับทราบการเรียก P80 30 นาที (ส.ค. 529 นาที) · เครื่องหยุด→เปิดใบซ่อม 3.2 นาที · ใบซ่อม→ซ่อมเสร็จ P80 0.9 ชม.\n"
     "ก่อนแก้ 21/09: เรียกช่างครั้งเดียวเด้งถึงช่างทั้ง 13 คน 278 ครั้ง/30 วัน — ตอนนี้เด้งเฉพาะทีมที่ถูกเรียก\n"
     "Demo: /daily-report กด 📞 เรียกช่าง → เลือกทีม → เปิด /mtn-repair ใบเดียวกันเห็นข้อมูลเครื่องเติมแล้ว → /mtn-analysis?tab=qc7\n"
     "ห้ามพูดเกิน: ซ่อมเร็วแล้ว แต่ปิดใบช้า — ใบค้างที่ขั้น QA 205 ใบ (06/10) · ห้ามใช้เลข \"ใบซ่อม 7 ใบ\" ใน infographic เก่า")

dept("JIG & DIE Maintenance", "Equipment records that answer 'where is it, what state?'", *T_NOT,
     ["JIG repair form FM-JIG-008 on paper",
      "Shim records not kept (customer request)",
      "Die location and state known by memory",
      "JIG team was called to DIE jobs"],
     ["Repair orders routed to the JIG or DIE team only",
      "Die register + storage map with pins + live status",
      "Open repair order turns the die red automatically",
      "Fixture shim record per point (customer requirement)"],
     ["Calls to the wrong team",
      "Searching for dies on the floor",
      "Separate paper shim log"],
     "receive calls from Production · report die / jig readiness back to Production",
     "JIG/DIE: ระบบพร้อมแล้ว แต่ยังไม่เริ่มใช้จริง — สถานะแม่พิมพ์ยังไม่เคยอัปเดต · ชิม fixture = 0 · ใบตรวจจิ๊กค้างรออนุมัติ\n"
     "สิ่งที่พูดได้จริง: ระบบชี้ว่าคิวกองอยู่ที่ทีมไหน (24/09: ใบรอทีม JIG รับ 23 ใบ นานสุด 22 วัน)\n"
     "Demo ใช้ได้เฉพาะถ้าทีม DIE วางหมุดบนผัง + อัปเดตสถานะจริงก่อนวันโชว์ (/equipment?tab=die) · ห้ามเปิดจอ fixture shim\n"
     "ระวัง: \"262 แม่พิมพ์\" (นับตัว) กับ \"92 ชุด\" (นับชุด) คนละหน่วย ให้เลือกพูดอย่างใดอย่างหนึ่ง")

dept("Store", "From reading the LINE chat to a pull board", *T_PART,
     ["Requests read from the LINE group chat",
      "\"First to ask gets served\" — no rounds, no phases",
      "Child-part demand calculated by hand",
      "Stock adjusted without approval"],
     ["Store board: WIP calls, child parts, purchasing, raw material, FG receipts",
      "Closing an FG order explodes the BOM into requests automatically",
      "Pick scan blocks the wrong part; delivery-point scan",
      "Goods counted before they enter stock"],
     ["Reading chat to know who needs what",
      "Calculating child-part demand",
      "Wrong-part deliveries (blocked by scan)"],
     "receives WIP calls from Production · sends purchase needs to Purchasing · receives finished goods from lines",
     "สโตร์: เล่าด้วยคำของทีมเอง (27/08) \"อ่านจาก LINE chat กลุ่มว่ามีใครเบิกงาน เห็นใครเบิกก่อนก็จัดของไปส่งเลย มั่วมาก\"\n"
     "ขาที่ระบบทำเองทำงานครบ (ปิดใบผลิต/ระเบิด BOM สำเร็จ 15,094/15,185 ครั้ง ณ 24/09) — ขาที่ต้องมีคนกดยังค้าง\n"
     "Demo: /heijunka แท็บคิวเติม WIP ให้เห็นใบที่เดินครบทุกขั้น หรือ /flow-tower ให้เห็นว่าปิดใบ FG ใบเดียวไหลไปถึงใบสั่งซื้อ\n"
     "ห้ามพูดเกิน: ไลน์ประกอบยังไม่ตั้ง min/max · ห้ามวาดว่ามีขั้นสแกน lot · ยอดรับของ bulk 01/10 (~2 ล้านชิ้น) ยังรอยืนยันว่าเป็นของจริง")

dept("Logistic / Delivery / Warehouse", "Customer mail in, delivery rounds out", *T_PART,
     ["Ford EDI 830/862 attachments saved and uploaded by hand",
      "e-SMART downloaded every ~2 hours",
      "FG stock kept separately from production"],
     ["EDI mail queued automatically every 15 min, one click to import",
      "e-SMART upload = customer's final confirmation",
      "Delivery countdown in 4 phases; 'shipped' deducts stock",
      "FG enters stock when the line closes the order"],
     ["Saving attachments by hand",
      "Re-keying FG receipts",
      "Missing which round has no stock (shown red)"],
     "receives customer demand · receives FG from Production · alerts Production when stock is short",
     "Logistic: Demo /planner-sales แผง 📬 ไฟล์จากเมลรอนำเข้า → กดนำเข้า → /customer-demand ดูสีรอบส่ง (เมลเข้า → แผนส่ง โดยไม่มีคนเปิดไฟล์แนบ)\n"
     "ใช้เลขใหม่: ปิดส่งลูกค้า 232 เที่ยว/8 วัน (30/09) — ห้ามใช้เลข \"38 จาก 472\" ของ 25/08\n"
     "ห้ามพูดเกิน: ยังใช้บัญชีร่วมเป็นหลัก · ตัวดึงเมลทำงานเฉพาะตอนเครื่อง Outlook เปิด และยังต้องกดยืนยันนำเข้า")

dept("Sales & Planning", "From a 13-sheet Excel to a live board", *T_PART,
     ["13-sheet monitoring Excel updated by hand every day",
      "Kanban calculated in Excel each month",
      "Customer part numbers matched to MAT by hand"],
     ["/monitoring: same sheets, IN / OUT / BALANCE filled from real orders and stock",
      "Kanban calculator uses the company calendar and master data",
      "Production plan: how many shifts, which days need OT",
      "Customer part → MAT matched automatically, mismatches flagged"],
     ["Typing production IN / OUT by hand",
      "Re-entering product data every month",
      "Silent part-number mismatches"],
     "receives customer demand · gives the plan frame to Production · gives lot sizes to Store",
     "Planning: Demo /monitoring เปิดบอร์ดชีทเดียวกับที่ทีมทำใน Excel แล้วชี้แถว IN ที่ระบบเติมเอง เทียบกับช่องที่ยังว่าง (บอกตรงๆ ว่าว่างเพราะพาร์ทนั้นยังไม่มีใบผลิตในระบบ)\n"
     "ห้ามพูดเกิน: วางแผนผลิต + คำนวณ kanban ยังไม่ถูกใช้จริง · มีประวัติผลิตในระบบแค่ 26/106 พาร์ทของชีทไลน์ปั๊ม")

dept("Quality (QA)", "Ready in the system — adoption starts now", *T_NOT,
     ["Suspect / scrap bin tags and logs on paper",
      "Writing material requests to Production for test samples",
      "Production writes scrap forms for parts QA took"],
     ["Sequential inspection; out-of-spec 'pass' is blocked",
      "Yellow / red bins with 4-way QA decision and tag age",
      "Material request FM-STO-003 → approve → Store issues",
      "QA-taken parts flow into the scrap report automatically"],
     ["Double paperwork between QA and Production",
      "Lost track of bins past their tag age"],
     "receives bins and 4M from Production · checks repairs after MTN (step 5) · decisions flow back to Production",
     "QA: ส่วนนี้ต้องพูดตรงที่สุด — ระบบพร้อม แต่ยังไม่เริ่มใช้จริง (ถังแดงค้าง 33 ใบ · QA ยังไม่เคยตัดสินถัง · ใบซ่อมรอ QA ตรวจ 205 ใบ)\n"
     "ข้อแนะนำ: อย่าเปิดจอ QA โชว์จนกว่าจะตัดสินถังแดงครบ — ถ้าตัดสินแล้ว ให้โชว์ /qa แท็บถัง → ใบที่ตัดสิน \"ทำลาย\" → ไปเห็นที่ใบ FM-PD2-002\n"
     "สิ่งที่ QA ควรประกาศบนเวที: แผนเคลียร์ใบรอตรวจ + เริ่มตัดสินถังทุกวัน")

dept("Process Engineering (PE)", "One data set instead of three Excel files", *T_PART,
     ["PFC, PFMEA and Control Plan in 3 separate Excel files",
      "Keeping operations consistent across files by hand",
      "Revision history in file names"],
     ["/pe-docs: one data set, three views on the same operations",
      "Import existing Excel; export PFMEA / CP to Excel",
      "Process flow drawn from the data",
      "Revision linked to its cause (claim / NCR / defect)"],
     ["Maintaining 3 files in sync",
      "Searching which revision came from which claim"],
     "shares PFMEA / CP with QA and Production · receives cycle times and claims",
     "PE: Demo /pe-docs ชุด P703 (Line 60) แท็บ Flow → ผังที่ระบบวาดเอง → แท็บ PFMEA เดียวกัน OP เดียวกัน (ข้อมูลจริงครบ) แล้วเชื่อมไป /order-trace ให้เห็นการสอบกลับถึง PFMEA\n"
     "ห้ามพูดเกิน: เอกสาร PE ไม่มีการอัปเดตตั้งแต่ 18 ส.ค. · PE ยังค้างส่ง CT 11 พาร์ท (ทำให้ %P ต่ำปลอม) · ห้ามเปิดจอ closed-loop 8D")

dept("HR & Administration", "Training, attendance and the calendar everyone depends on", *T_PART,
     ["OJT form FM-HRM-004 signed on paper",
      "Monthly KPI pack (FM-HRM-6-022/024/025) printed and signed",
      "Attendance and OT bus booking by hand"],
     ["/ojt-training: scores and signatures on screen, printable",
      "KPI pack exported to Excel in the HR form layout",
      "Check-in: attendance, PPE and OT bus in one screen",
      "Company calendar drives kanban, plan, PM and LPA"],
     ["Printing and chasing signatures",
      "Re-typing KPI packs",
      "Wrong working-day counts (calendar now set to 2027)"],
     "the calendar feeds every department · attendance feeds manpower and skills",
     "HR/ธุรการ: เน้นว่าปฏิทินบริษัทเป็น \"ข้อมูลกลาง\" — ถ้าผิด kanban / แผนผลิต / PM / LPA ผิดตามหมดเงียบๆ (บทเรียนจริง 28–29 ก.ย.) · ตอนนี้ปฏิทินมีถึงปี 2027 แล้ว\n"
     "ก่อนพูด: ยังไม่มีตัวเลขการใช้งาน OJT จริงในเอกสาร ต้องคิวรีก่อน (ถ้าไม่มีใบจริง อย่าเปิดจอ)")

dept("Management / OBEYA", "From a paper wall to one scoreboard", *T_PART,
     ["Paper OBEYA KPI board on the wall",
      "Review deck and Excel even disagreed on targets",
      "Decisions stayed in heads, chat or paper",
      "Tracing = searching files, asking shift leaders"],
     ["/obeya: monthly KPI + daily SQDCM, one scoring rule for every screen",
      "Factory map and TV Andon for every line",
      "/order-trace: one order number shows the whole chain",
      "Action board: owner and due date for every miss"],
     ["Arguing which number is right",
      "Updating the wall board by hand",
      "Hours of tracing (seconds now)"],
     "sees every department's data · assigns actions back to departments",
     "ผู้บริหาร: Demo เส้นทางที่ล็อกไว้ /factory-map → /tv → /obeya → ไลน์ 060/061 → /order-trace → /mtn-repair → /program-update\n"
     "ห้ามพูดเกิน: ACTION BOARD ยังไม่มีใบจริง (ต้องมี 5–10 ใบก่อนวันโชว์) · /group-overview เป็น mockup · ยังไม่ต่อ SAP/SCADA · ยังไม่ย้ายลง server บริษัท")

# ── Honest status ───────────────────────────────────────────────────
s = content("What is still slow — and we say it openly", MEASURED)
stat(s, 0.5, 1.85, 2.38, "205", "Repair orders waiting for QA", "median wait 13 days")
stat(s, 3.0, 1.85, 2.38, "97 h", "Repair order → closed", "median · Sep (fixed fast, closed slow)")
stat(s, 5.5, 1.85, 2.38, "13 h", "Shift close → approved", "median · Sep (Jul: 4.7 h)")
stat(s, 8.0, 1.85, 2.38, "0", "QA decisions on scrap bins", "33 red bins open")
stat(s, 10.5, 1.85, 2.33, "0", "Action items from meetings", "since 13 Jul")
T.headline_box(s, 0.5, 3.9, 12.33, "Pattern: steps the system does by itself work — steps that need a person to press are where work waits")
for i, t in enumerate([
    "QA: daily session to clear repair orders at step 5 and decide every bin",
    "Supervisors: approve shift-close requests the same day",
    "Every morning meeting: at least one action item with an owner and a due date",
]):
    T.text(s, 0.7, 4.45 + i * 0.55, 12.0, 0.5, [[("➜  " + t, {"sz": 13, "b": True, "c": GREEN_M})]], anchor=MSO_ANCHOR.MIDDLE)
notes(s, "สไลด์นี้ต้องพูด — GM ต่างโรงจะถามแน่ว่า \"ใช้จริงแค่ไหน\" การบอกเองดีกว่าโดนจับได้\n"
          "สรุปเดียว: ส่วนที่ระบบทำเองทำงานครบ · ส่วนที่ค้างคือขั้นที่ต้องมีคนกด ⇒ แต่ละส่วนงานรับปากเรื่องไหนบนเวที")

# ── Next steps ─────────────────────────────────────────────────────
s = content("Next steps to the roadshow", "Owner per department · 06 – 30 Oct 2026")
steps = [("Week 2  (12–16 Oct)", "QA clears bins & step-5 repairs · DIE updates die status · first real action items"),
         ("Week 3  (19–22 Oct)", "Each department rehearses its own 3 minutes · fill in the 'before' numbers · collect 3 real cases"),
         ("Week 4  (26–30 Oct)", "Feature freeze · final numbers measured the morning of the show · roadshow")]
for i, (a, b) in enumerate(steps):
    yy = 1.9 + i * 1.45
    T.rect(s, 0.5, yy, 3.3, 1.15, fill=GREEN_D, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, 0.5, yy, 3.3, 1.15, [[(a, {"sz": 16, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    T.rect(s, 3.95, yy, 8.88, 1.15, fill=GREEN_TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, 4.15, yy, 8.5, 1.15, [[(b, {"sz": 14, "c": GREEN_M})]], anchor=MSO_ANCHOR.MIDDLE)
notes(s, "อ้างอิงไทม์ไลน์คุมงานใน docs/ADOPTION-ROADMAP.md · ตัวเลขบนสไลด์ทั้งหมดจะวัดใหม่เช้าวันโชว์ และเขียนวันที่วัดไว้ทุกสไลด์")

# ── Closing ─────────────────────────────────────────────────────────
s = T.blank(p)
T.rect(s, 0, 0, T.SW, T.SH, fill=GREEN_D)
T.text(s, 1.0, 2.6, 11.3, 1.0, [[("Before We Build Parts, We Build People", {"sz": 34, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
T.text(s, 1.0, 3.8, 11.3, 0.8, [[("Thank you", {"sz": 28, "b": True, "c": GOLD})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "ปิด: ระบบไม่ได้ทำงานแทนคน — ระบบทำให้คนทุกแผนกเห็นข้อมูลชุดเดียวกัน และส่งงานต่อกันได้เร็วขึ้น")

p.save(OUT)
print("saved", os.path.abspath(OUT), "slides:", len(p.slides))

# -*- coding: utf-8 -*-
"""ขยาย ESM ไปโรงงานใหม่ (ฐานข้อมูลแยก) — ต้องเตรียมอะไร · ธีม TSG R01 · ภาษาไทย · ฉบับอ่านง่าย (2026-10-28)
Run:  python3 build_new_factory_deploy.py  →  ../ESM_New_Factory_Deploy_Checklist_2026-10-28.pptx
ผู้ฟัง = ผู้บริหาร / หัวหน้าส่วนงานของโรงงานใหม่ ⇒ ภาษาบ้านๆ + ภาพ/ไอคอนแทนตัวหนังสือ
ศัพท์เทคนิคทั้งหมดรวมไว้สไลด์ "ภาคผนวก สำหรับทีม IT" ที่เดียว
ตัวเลขนับจากรีโป 28/10/2026: migration 707 ไฟล์ · edge function 24 · ที่เก็บรูป 15 ชุด ·
URL ฐานข้อมูล TSAT ฝังในโค้ด 5 หน้า (Checkin/DailyReport/MtnRepair/PMCheckData/PMSchedule) + DR fallback
ค่า Supabase อ้างอิงยอดชำระจริง 11 ต.ค. = $36.70 (docs/modules/deploy.md)
"""
import os
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import tsg_theme as T
from tsg_theme import GREEN, GREEN_D, GREEN_TINT, BORDER, ORANGE, ORANGE_LT, AMBER, GREY, WHITE

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "ESM_New_Factory_Deploy_Checklist_2026-10-28.pptx")
EMOJI = "Segoe UI Emoji"
p = T.new_pres()
page = [0]


def notes(s, th):
    s.notes_slide.notes_text_frame.text = th


def content(title, sub=None):
    s = T.blank(p); page[0] += 1
    T.head(s, title, sub)
    T.footer(s, page[0])
    return s


def icon(s, x, y, d, emo, fill=GREEN_TINT, line=None):
    """ไอคอนวงกลม + emoji ตรงกลาง (ภาพแทนตัวหนังสือ)"""
    T.rect(s, x, y, d, d, fill=fill, line=line, lw=2, shape=MSO_SHAPE.OVAL)
    T.text(s, x, y, d, d, [[(emo, {"sz": int(d * 30), "f": EMOJI, "c": GREEN})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def label(s, x, y, w, big, small=None, c=GREEN, bsz=17, ssz=13, align=PP_ALIGN.CENTER):
    runs = [[(big, {"sz": bsz, "b": True, "c": c})]]
    if small:
        runs.append([(small, {"sz": ssz, "c": GREY})])
    T.text(s, x, y, w, 1.0, runs, align=align, sa=3)


def tile(s, x, y, w, h, emo, big, small=None, fill=WHITE, accent=GREEN_D):
    T.rect(s, x, y, w, h, fill=fill, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    T.rect(s, x, y, w, 0.1, fill=accent)
    d = min(1.05, h * 0.42)
    icon(s, x + (w - d) / 2, y + 0.28, d, emo)
    label(s, x + 0.1, y + 0.4 + d, w - 0.2, big, small, bsz=16, ssz=12.5)


def arrow(s, x, y, w=0.6, h=0.5, c=ORANGE):
    T.rect(s, x, y, w, h, fill=c, shape=MSO_SHAPE.RIGHT_ARROW)


# ── 1 ปก ──
s = T.title_slide(p, "ขยาย ESM ไปโรงงานใหม่ — ต้องเตรียมอะไรบ้าง",
                  "โปรแกรมชุดเดียวกัน · ข้อมูลของใครของมัน",
                  "Enterprise Shopfloor Management (ESM)", "Thai Summit Group · ตุลาคม 2569")
page[0] += 1
notes(s, "เปิด: ถ้าโรงงานอื่นในกลุ่มอยากใช้ ESM ต้องเตรียมอะไร — คำตอบสั้นๆ คือ 3 ส่วน: ที่ตั้งระบบ · ปรับโปรแกรมให้เป็นของโรงงานนั้น · ข้อมูลของโรงงาน\n"
         "ส่วนที่นานที่สุดไม่ใช่ฝั่ง IT แต่เป็นข้อมูลจากโรงงาน")

# ── 2 ภาพเปรียบเทียบ: บ้านแบบเดียวกัน ตู้เซฟคนละใบ ──
s = content("เหมือนสร้างบ้านหลังใหม่ จากแบบบ้านเดิม", "ใช้แบบเดียวกัน แต่ของในบ้าน กุญแจ และตู้เซฟ เป็นของเจ้าของบ้านแต่ละหลัง")
for x, name, fill in [(0.9, "TSAT (ใช้อยู่)", GREEN_TINT), (8.43, "โรงงานใหม่", ORANGE_LT)]:
    T.rect(s, x, 1.85, 4.0, 3.3, fill=fill, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, x, 1.95, 4.0, 0.5, [[(name, {"sz": 20, "b": True, "c": GREEN_D})]], align=PP_ALIGN.CENTER)
    for i, (emo, cap) in enumerate([("🗄️", "ตู้เซฟข้อมูล"), ("🔑", "กุญแจ / ผู้ใช้"), ("👷", "คน · ไลน์ · เครื่อง")]):
        xx = x + 0.25 + i * 1.25
        icon(s, xx, 2.65, 1.0, emo, fill=WHITE)
        T.text(s, xx - 0.15, 3.75, 1.3, 0.8, [[(cap, {"sz": 13, "b": True, "c": GREEN})]], align=PP_ALIGN.CENTER)
icon(s, 5.67, 1.95, 2.0, "📐", fill=WHITE, line=ORANGE)
T.text(s, 5.2, 4.05, 2.93, 0.9, [[("แบบบ้าน = โปรแกรม ESM", {"sz": 15, "b": True, "c": ORANGE})], [("ใช้ร่วมกัน แก้ที่เดียว ได้ทุกโรง", {"sz": 12.5, "c": GREY})]],
       align=PP_ALIGN.CENTER)
arrow(s, 5.05, 2.7, 0.55, 0.5, GREEN); arrow(s, 7.75, 2.7, 0.55, 0.5, ORANGE)
T.headline_box(s, 0.9, 5.5, 11.53, "ข้อมูลไม่ปนกัน · คนโรงงานหนึ่งมองไม่เห็นข้อมูลอีกโรงงาน")
notes(s, "เปรียบเทียบให้เห็นภาพ: โปรแกรมเป็นแบบบ้าน ใช้ร่วมกันได้ — ปรับปรุงที่ TSAT แล้วโรงงานอื่นได้ด้วย\n"
         "แต่ตู้เซฟ (ฐานข้อมูล) กุญแจ (บัญชีผู้ใช้) และของในบ้าน (คน ไลน์ เครื่องจักร สินค้า) ต้องแยกกันเด็ดขาด")

# ── 3 3 ส่วนที่ต้องเตรียม + ใครทำ ──
s = content("ต้องเตรียม 3 ส่วน", "แต่ละส่วนมีเจ้าภาพชัดเจน")
parts = [("🖥️", "① ที่ตั้งระบบ", "เจ้าภาพ: IT", "พื้นที่เก็บข้อมูล · ที่อยู่เว็บ · เข้าจากมือถือได้", GREEN_D),
         ("🛠️", "② ปรับโปรแกรม", "เจ้าภาพ: ทีม ESM", "เปลี่ยนชื่อ โลโก้ ฟอร์ม ให้เป็นของโรงงานใหม่", GREEN),
         ("📋", "③ ข้อมูลโรงงาน", "เจ้าภาพ: โรงงานใหม่", "ผังองค์กร · คน · ไลน์ · เครื่องจักร · สินค้า", ORANGE)]
for i, (emo, h, who, d, col) in enumerate(parts):
    x = 0.7 + i * 4.1
    T.rect(s, x, 1.75, 3.75, 4.3, fill=WHITE, line=col, lw=2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    icon(s, x + 1.12, 2.0, 1.5, emo, fill=GREEN_TINT)
    T.text(s, x, 3.65, 3.75, 0.5, [[(h, {"sz": 22, "b": True, "c": col})]], align=PP_ALIGN.CENTER)
    T.rect(s, x + 0.8, 4.2, 2.15, 0.42, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, x + 0.8, 4.2, 2.15, 0.42, [[(who, {"sz": 13, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    T.text(s, x + 0.25, 4.8, 3.25, 1.1, [[(d, {"sz": 14, "c": GREEN})]], align=PP_ALIGN.CENTER)
T.text(s, 0.7, 6.2, 12, 0.45, [[("⏱  ส่วนที่ ③ ใช้เวลานานที่สุด — ควรเริ่มเก็บข้อมูลตั้งแต่วันแรก", {"sz": 15, "b": True, "c": ORANGE})]], align=PP_ALIGN.CENTER)
notes(s, "3 ส่วน 3 เจ้าภาพ — IT เตรียมที่ตั้งระบบ, ทีม ESM ปรับโปรแกรม, โรงงานใหม่เตรียมข้อมูลของตัวเอง\n"
         "ส่วนที่ 1-2 ทำได้ในไม่กี่วัน ส่วนที่ 3 ขึ้นกับความพร้อมข้อมูลของโรงงาน")

# ── 4 ส่วน 1 ที่ตั้งระบบ ──
s = content("① ที่ตั้งระบบ — IT เตรียม", "เปิดบัญชีใหม่ทั้งหมด ไม่ใช้ของ TSAT ร่วม")
tiles = [("🗄️", "ตู้เก็บข้อมูล 2 ใบ", "ข้อมูลคน · ข้อมูลผลิต (แยกกันเหมือน TSAT)"),
         ("🌐", "ที่อยู่เว็บใหม่", "เช่น esm-ชื่อโรงงาน + ใบรับรองความปลอดภัย"),
         ("📱", "เข้าได้จากมือถือ", "ทั้งในโรงงานและนอกโรงงาน (4G/5G)"),
         ("🔔", "ช่องแจ้งเตือน", "แจ้งในแอปเป็นหลัก · Telegram ถ้าต้องการ"),
         ("💾", "สำรองข้อมูล", "ตั้งสำรองอัตโนมัติ + ซ้อมกู้คืน"),
         ("💰", "ค่าใช้จ่าย ≈ $37 / เดือน", "อ้างอิงที่ TSAT จ่ายจริง (ค่าที่เก็บข้อมูล)")]
for i, (emo, b, sm) in enumerate(tiles):
    x = 0.55 + (i % 3) * 4.15; y = 1.65 + (i // 3) * 2.55
    tile(s, x, y, 3.9, 2.35, emo, b, sm, accent=ORANGE if emo == "💰" else GREEN_D)
notes(s, "ฝั่ง IT: เปิดที่เก็บข้อมูลใหม่ 2 ชุด (แบบเดียวกับ TSAT ห้ามรวมเป็นชุดเดียว — ชุดข้อมูลผลิตเปิดอ่านกว้างกว่า)\n"
         "ค่าใช้จ่ายอ้างอิงยอดที่ TSAT ชำระจริง 11 ต.ค. = $36.70/เดือน · ที่วางหน้าเว็บใช้ของฟรีหรือ server บริษัทได้\n"
         "รายละเอียดเทคนิคอยู่ภาคผนวกท้ายเด็ค")

# ── 5 ส่วน 2 ปรับโปรแกรม ──
s = content("② ปรับโปรแกรม — ถอดป้าย TSAT ออก ติดป้ายโรงงานใหม่", "ทีม ESM ทำ · ต้องเสร็จก่อนเปิดใช้ ไม่งั้นข้อมูลไหลผิดโรง")
items = [("🏷️", "ชื่อ + โลโก้", "บนจอ ไอคอนแอป ใบพิมพ์"),
         ("📄", "เลขฟอร์มเอกสาร", "เลข FM / Rev ของโรงงานใหม่"),
         ("📨", "ปลายทางแจ้งเตือน", "ให้วิ่งเข้าโรงงานใหม่ ไม่ใช่ TSAT"),
         ("⚙️", "ค่าที่เป็นของ TSAT", "เช่น รหัสคลัง · ไลน์ตัวอย่าง")]
for i, (emo, b, sm) in enumerate(items):
    x = 0.55 + i * 3.12
    T.rect(s, x, 1.8, 2.9, 3.4, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    icon(s, x + 0.85, 2.05, 1.2, emo)
    T.rect(s, x + 0.35, 3.45, 0.95, 0.4, fill=GREY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, x + 0.35, 3.45, 0.95, 0.4, [[("TSAT", {"sz": 12, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    arrow(s, x + 1.33, 3.5, 0.3, 0.3)
    T.rect(s, x + 1.66, 3.45, 0.95, 0.4, fill=ORANGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, x + 1.66, 3.45, 0.95, 0.4, [[("ใหม่", {"sz": 12, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    label(s, x + 0.1, 4.0, 2.7, b, sm, bsz=16, ssz=12.5)
T.headline_box(s, 0.55, 5.55, 12.25, "⚠  ปลายทางแจ้งเตือนสำคัญที่สุด — ตอนนี้ยังชี้มาที่ TSAT ใน 5 หน้าจอ ต้องแก้ก่อน")
notes(s, "สิ่งที่ฝังเป็นของ TSAT อยู่ในโปรแกรม: ชื่อ/โลโก้, เลขฟอร์ม, และที่สำคัญคือปลายทางแจ้งเตือน\n"
         "ตอนนี้มี 5 หน้าจอ (เช็คชื่อ · รายงานผลิต · ใบแจ้งซ่อม · ใบตรวจ PM · แผน PM) ที่ส่งแจ้งเตือนตรงเข้าระบบ TSAT\n"
         "ถ้าไม่แก้ก่อน โรงงานใหม่กดแจ้งซ่อม แจ้งเตือนจะเด้งมาที่ TSAT — ทีม ESM แก้ให้อ่านจากการตั้งค่าแทน ทำครั้งเดียวใช้ได้ทุกโรง")

# ── 6 divider ──
T.section_divider(p, "divider_r01.jpg", ["③ ข้อมูลที่โรงงานใหม่", "ต้องเตรียม"])
page[0] += 1

# ── 7 บันไดข้อมูล: ลงจากฐานขึ้นบน ──
s = content("ข้อมูลที่ต้องมีก่อนวันแรก — ลงจากขั้นล่างขึ้นบน", "ขั้นล่างเป็นฐานของขั้นบน ข้ามขั้นไม่ได้ · ส่งเป็น Excel ตามแม่แบบที่ทีม ESM ให้")
steps = [("🏢", "ผังองค์กร", "ฝ่าย · แผนก · ทีม · Cost Center", "HR / บัญชี"),
         ("🏭", "ไลน์ผลิต", "ชื่อไลน์ · คนมาตรฐานต่อกะ", "ฝ่ายผลิต"),
         ("👷", "พนักงาน", "รหัส · ชื่อ · ทีม · รูปถ่าย", "HR"),
         ("🔑", "ผู้ใช้ระบบ", "ใครใช้ · ดูแลส่วนไหน", "หัวหน้าส่วน"),
         ("📅", "ปฏิทิน + กะ", "วันหยุด · OT · เวลาพัก", "HR / ผลิต"),
         ("⚙️", "เครื่องจักร", "เลขเครื่อง · อยู่ไลน์ไหน", "ซ่อมบำรุง"),
         ("📦", "สินค้า", "รหัส · ลูกค้า · เวลาผลิตต่อครั้ง", "PE / วางแผน")]
n = len(steps); w = 1.72
for i, (emo, b, sm, who) in enumerate(steps):
    x = 0.45 + i * (w + 0.05)
    y = 5.0 - i * 0.48
    h = 6.55 - y
    T.rect(s, x, y, w, h, fill=GREEN_TINT if i % 2 == 0 else WHITE, line=BORDER, shape=MSO_SHAPE.RECTANGLE)
    T.rect(s, x, y, w, 0.08, fill=ORANGE if i == 0 else GREEN_D)
    icon(s, x + 0.41, y - 1.0, 0.9, emo, fill=WHITE, line=GREEN)
    T.text(s, x + 0.05, y + 0.15, w - 0.1, 0.4, [[(f"{i + 1}. {b}", {"sz": 15, "b": True, "c": GREEN_D})]], align=PP_ALIGN.CENTER)
    T.text(s, x + 0.08, y + 0.6, w - 0.16, 0.8, [[(sm, {"sz": 12, "c": GREEN})]], align=PP_ALIGN.CENTER)
    T.text(s, x + 0.05, 6.15, w - 0.1, 0.35, [[(who, {"sz": 12, "b": True, "c": ORANGE})]], align=PP_ALIGN.CENTER)
notes(s, "ลำดับสำคัญ: ไม่มีผังองค์กร = ผูกไลน์ไม่ได้ · ไม่มีไลน์ = ผูกพนักงานไม่ได้ · ไม่มีพนักงาน = สร้างผู้ใช้ไม่ได้\n"
         "ชื่อเจ้าของข้อมูลด้านล่างแต่ละขั้น = คนที่ต้องยืนยันว่าข้อมูลถูก (ลงชื่อ + วันที่)\n"
         "\"เวลาผลิตต่อครั้ง\" = เวลาปั๊ม/ผลิต 1 จังหวะ (ไม่ใช่ต่อชิ้น) — งานคู่ซ้าย/ขวาต้องบอกด้วยว่าเป็นคู่")

# ── 8 เติมทีหลัง ──
s = content("ข้อมูลที่เติมทีหลัง — เปิดเรื่องไหน เตรียมเรื่องนั้น", "ไม่ต้องครบวันแรก · เริ่มที่ไลน์นำร่อง 1–2 ไลน์ให้ใช้ได้จริงก่อน แล้วค่อยขยาย")
mods = [("🧩", "BOM / ขั้นตอนผลิต", "PE"), ("🚚", "ลูกค้า / ผู้ส่งของ", "Logistic"), ("🔩", "แม่พิมพ์ · จิ๊ก · อะไหล่", "DIE / JIG / สโตร์"),
        ("🔧", "แผนซ่อมบำรุง (PM)", "ซ่อมบำรุง"), ("✅", "ตรวจคุณภาพ", "QA"), ("📊", "KPI / เป้าหมาย", "ผู้บริหาร / บัญชี"),
        ("🦺", "PPE / ความปลอดภัย", "Safety"), ("🗂️", "ทะเบียนเอกสาร", "Document Control")]
for i, (emo, b, who) in enumerate(mods):
    x = 0.55 + (i % 4) * 3.12; y = 1.65 + (i // 4) * 2.5
    tile(s, x, y, 2.9, 2.3, emo, b, who)
notes(s, "เรื่องพวกนี้เปิดทีละโมดูลตามความพร้อม ไม่ต้องรอครบทั้งหมด\n"
         "หลักที่ใช้ที่ TSAT: ทำไลน์ตัวอย่าง 1-2 ไลน์ให้ข้อมูลเชื่อมกันครบทุกเรื่องก่อน ดีกว่ามีข้อมูลกว้างแต่ไม่เชื่อมกัน")

# ── 9 ทำ/ไม่ทำ ──
s = content("เตรียมข้อมูลให้ถูกตั้งแต่แรก — ไม่ต้องตามแก้ทีหลัง", "ทุกข้อเคยเกิดจริงที่ TSAT")
rules = [("เลขเครื่องเขียนแบบเดียว", "SW10", "SW-10 · sw 10"),
         ("ไม่รู้ = เว้นว่าง", "(ว่าง)", "0 · “อื่นๆ”"),
         ("เวลาผลิต = ต่อ 1 ครั้ง", "12 วิ / ครั้ง", "6 วิ / ชิ้น"),
         ("ข้อมูลของโรงงานเอง", "ของโรงงานใหม่", "ก๊อปจาก TSAT")]
for i, (h, ok, ng) in enumerate(rules):
    x = 0.55 + i * 3.12
    T.text(s, x, 1.7, 2.9, 0.5, [[(h, {"sz": 16, "b": True, "c": GREEN_D})]], align=PP_ALIGN.CENTER)
    for j, (emo, val, col, fill) in enumerate([("✅", ok, GREEN, GREEN_TINT), ("❌", ng, ORANGE, ORANGE_LT)]):
        y = 2.35 + j * 1.85
        T.rect(s, x, y, 2.9, 1.65, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        T.text(s, x, y + 0.1, 2.9, 0.75, [[(emo, {"sz": 30, "f": EMOJI, "c": col})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        T.text(s, x + 0.1, y + 0.85, 2.7, 0.7, [[(val, {"sz": 18, "b": True, "c": col})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
T.text(s, 0.55, 6.15, 12.3, 0.45, [[("🔒  ข้อมูลพนักงานและรูปถ่าย ใช้ตามนโยบายข้อมูลส่วนบุคคล (PDPA) ของโรงงานนั้น", {"sz": 14, "b": True, "c": GREEN})]],
       align=PP_ALIGN.CENTER)
notes(s, "ตัวอย่างจริงที่ TSAT: เลขเครื่องเขียน 2 แบบ ทำให้กราฟแตกเป็น 2 แท่งทั้งที่เป็นเครื่องเดียวกัน\n"
         "ใส่ 0 แทนการเว้นว่าง ทำให้ระบบเข้าใจว่า 'รู้แล้วว่าเป็นศูนย์' ตัวเลข OEE เพี้ยน\n"
         "งานปั๊มคู่ซ้าย/ขวา ปั๊มครั้งเดียวได้ 2 ชิ้น — ถ้าให้เวลาต่อชิ้น เวลามาตรฐานจะผิด")

# ── 10 ข้อห้าม ──
s = content("4 เรื่องที่ห้ามทำ", "พลาดแล้วแก้ยาก")
donts = [("🗄️", "ใช้ที่เก็บข้อมูลร่วมกับ TSAT", "ข้อมูลสองโรงปนกันทันที"),
         ("🔔", "เปิดใช้ก่อนแก้ปลายทางแจ้งเตือน", "แจ้งเตือนโรงงานใหม่เด้งมาที่ TSAT"),
         ("🍴", "แยกโปรแกรมออกเป็นคนละชุด", "แก้ปัญหาที่ TSAT แล้ว โรงใหม่ไม่ได้ด้วย"),
         ("⏰", "อัปเดตระบบในเวลาทำงาน", "จอที่เปิดค้างพังพร้อมกันทุกเครื่อง")]
for i, (emo, b, why) in enumerate(donts):
    x = 0.55 + (i % 2) * 6.2; y = 1.7 + (i // 2) * 2.4
    T.rect(s, x, y, 5.95, 2.15, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    icon(s, x + 0.3, y + 0.45, 1.25, emo, fill=ORANGE_LT)
    T.rect(s, x + 1.2, y + 0.35, 0.45, 0.45, fill=ORANGE, shape=MSO_SHAPE.OVAL)
    T.text(s, x + 1.2, y + 0.35, 0.45, 0.45, [[("✕", {"sz": 16, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    T.text(s, x + 1.85, y + 0.4, 3.95, 1.5, [[(b, {"sz": 17, "b": True, "c": GREEN_D})], [(why, {"sz": 14, "c": GREY})]],
           anchor=MSO_ANCHOR.MIDDLE, sa=6)
notes(s, "(1) ห้ามใช้ฐานข้อมูลร่วม — ระบบไม่มีกำแพงกั้นระหว่างโรงงานในข้อมูลชุดเดียว\n"
         "(2) ปลายทางแจ้งเตือนต้องแก้ก่อนเปิดใช้ (3) ใช้โปรแกรมชุดเดียว ปรับด้วยการตั้งค่า ไม่แตกเป็นคนละชุด\n"
         "(4) อัปเดตหลังเลิกกะ — บทเรียนจริง 26/10 ที่ TSAT")

# ── 11 แผน 5 สัปดาห์ ──
s = content("แผนเปิดใช้ 5 สัปดาห์ (ข้อเสนอ)", "เริ่มจากไลน์นำร่อง แล้วขยายทีละกลุ่ม")
plan = [("🖥️", "สัปดาห์ 1", "เตรียมระบบ", "IT + ทีม ESM"),
        ("🛠️", "สัปดาห์ 2", "ปรับโปรแกรม", "ทีม ESM"),
        ("📋", "สัปดาห์ 3", "ลงข้อมูลหลัก", "โรงงานใหม่"),
        ("🚩", "สัปดาห์ 4", "ไลน์นำร่อง", "ใช้จริง 1–2 ไลน์"),
        ("🚀", "สัปดาห์ 5", "ขยายทั้งโรง", "เปิด PM / QA / KPI")]
for i, (emo, wk, h, who) in enumerate(plan):
    x = 0.45 + i * 2.5
    col = ORANGE if i == 3 else GREEN_D
    icon(s, x + 0.55, 1.75, 1.3, emo, fill=GREEN_TINT if i != 3 else ORANGE_LT)
    T.rect(s, x, 3.3, 2.4, 0.7, fill=col, shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
    T.text(s, x + 0.25, 3.3, 2.0, 0.7, [[(wk, {"sz": 15, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    T.text(s, x, 4.15, 2.4, 1.0, [[(h, {"sz": 18, "b": True, "c": col})], [(who, {"sz": 13, "c": GREY})]], align=PP_ALIGN.CENTER, sa=4)
T.headline_box(s, 0.45, 5.6, 12.4, "ส่งแม่แบบ Excel ให้โรงงานใหม่ตั้งแต่สัปดาห์ที่ 1 — ข้อมูลคือตัวกำหนดวันเปิดใช้")
notes(s, "5 สัปดาห์เป็นข้อเสนอเบื้องต้น ถ้าข้อมูลโรงงานพร้อมเร็ว เปิดไลน์นำร่องได้เร็วขึ้น\n"
         "สัปดาห์ 1-2 ทำขนานกับการเก็บข้อมูลได้ — อย่ารอให้ระบบเสร็จก่อนค่อยเริ่มเก็บข้อมูล")

# ── 12 ภาคผนวก IT ──
s = content("ภาคผนวก — รายละเอียดสำหรับทีม IT", "หน้าเดียวที่มีศัพท์เทคนิค · นับจากโค้ดจริง 28 ต.ค. 2569")
cols = [("ที่ตั้งระบบ", ["Supabase 2 projects: MAIN + DR (ห้ามรวม — DR เปิด anon)",
                         "Static hosting: Render หรือ nginx/IIS + TLS",
                         "Auth: ปิด sign-up · สร้างผู้ใช้ผ่าน /add-user",
                         "Backup / PITR + ซ้อมกู้คืน 1 ครั้ง"]),
        ("ฐานข้อมูล", ["Schema-only dump จาก MAIN / DR (migration 707 ไฟล์ ติดป้ายฝั่งไม่ครบ)",
                       "RLS + has_perm() · seed role_permissions",
                       "Realtime publication · pg_cron · Storage 15 buckets",
                       "Edge functions 24 ตัว + secrets ชุดใหม่"]),
        ("โค้ดที่ต้องแก้", ["URL send-notification ฝังใน 5 หน้า → อ่าน env",
                           "supabaseClient: DR fallback → ไม่มี env = หยุด",
                           "SchemaMap + edge fn 5 ตัว มี project id",
                           "ชื่อ/โลโก้/manifest → config แบรนด์ต่อโรงงาน"])]
for i, (h, ls) in enumerate(cols):
    x = 0.45 + i * 4.18
    T.rect(s, x, 1.7, 3.98, 3.6, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.rect(s, x, 1.7, 3.98, 0.46, fill=GREEN_D)
    T.text(s, x, 1.7, 3.98, 0.46, [[(h, {"sz": 15, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    T.text(s, x + 0.15, 2.3, 3.7, 4.1, [[("•  " + l, {"sz": 13.5, "c": GREEN})] for l in ls], sa=10)
notes(s, "สำหรับทีม IT เท่านั้น — ไม่ต้องนำเสนอในที่ประชุมทั่วไป\n"
         "env ฝั่งเว็บ: VITE_SUPABASE_URL / _ANON_KEY / _DR_URL / _DR_KEY · edge secrets: DR_URL, DR_ANON_KEY, CRON_SECRET, INGEST_SECRET, CLEANUP_TOKEN, TELEGRAM_* (ถ้าใช้), ANTHROPIC_API_KEY (ถ้าใช้)\n"
         "ทุก key ต้องออกใหม่ ห้ามใช้ของ TSAT ซ้ำ · รายละเอียด deploy: docs/modules/deploy.md")

# ── 13 ปิด ──
page[0] += 1
T.closing_slide(p, page[0], "Before We Build Parts, We Build People", [
    ("โปรแกรมเดียวกัน · ข้อมูลของใครของมัน", ORANGE),
    ("ขอบคุณ  ·  ถาม-ตอบ", GREEN)])

p.save(OUT)
print("saved", os.path.relpath(OUT), "slides:", len(p.slides))

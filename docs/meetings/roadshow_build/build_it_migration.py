# -*- coding: utf-8 -*-
"""สไลด์คุยกับไอที: ย้าย ESM เข้าเซิร์ฟเวอร์บริษัท (ธีม TSG R01 · ภาษาไทย · 2026-10-20)
Run:  python3 build_it_migration.py  →  ../TSAT_ESM_IT_Server_Migration_2026-10-20.pptx
ตัวเลข: นับจาก repo (edge functions/cron/migrations) + ขนาด DB/storage วัดผ่าน MCP 12/10
"""
import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
import tsg_theme as T
from tsg_theme import GREEN, GREEN_D, GREEN_TINT, BORDER, ORANGE, GREY, WHITE

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "TSAT_ESM_IT_Server_Migration_2026-10-20.pptx")
p = T.new_pres()
page = [0]


def notes(s, th):
    s.notes_slide.notes_text_frame.text = th


def content(title, sub=None):
    s = T.blank(p); page[0] += 1
    T.head(s, title, sub)
    T.footer(s, page[0])
    return s


def node(s, x, y, w, h, head, sub, fill=WHITE, hc=GREEN_D, line=BORDER):
    T.rect(s, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    T.text(s, x + 0.1, y + 0.08, w - 0.2, h - 0.16,
           [[(head, {"sz": 15, "b": True, "c": hc})]] + [[(l, {"sz": 12, "c": GREY})] for l in sub],
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sa=2)


def arrow(s, x1, y1, x2, y2, c=GREEN):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = c; ln.line.width = Pt(2)
    ln.line._get_or_add_ln().append(ln.line._ln.makeelement(
        '{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd', {'type': 'triangle'}))


# ── 1 ปก ──
s = T.title_slide(p, "ย้ายระบบ ESM เข้าเซิร์ฟเวอร์บริษัท",
                  "ระบบตอนนี้เป็นอย่างไร · มีทางเลือกอะไร · ต้องเตรียมอะไร",
                  "ประชุมร่วมทีมไอที · 20 ตุลาคม 2569 · 15:00",
                  "TSAT · Thai Summit Group",
                  decl_label="คำประกาศคุณภาพของฉัน")
page[0] += 1
notes(s, "ESM ตอนนี้เป็นเว็บแอปที่อยู่บนคลาวด์ 2 เจ้า หน้าเว็บอยู่ Render ฐานข้อมูลกับหลังบ้านอยู่ Supabase (PostgreSQL บนคลาวด์) อยากย้ายเข้าเซิร์ฟเวอร์บริษัท วันนี้ขอให้ไอทีช่วยเลือกวิธี และบอกว่าต้องเตรียมอะไร")

# ── 2 ระบบตอนนี้ ──
s = content("ระบบ ESM ตอนนี้ประกอบด้วยอะไร", "ภาษาหน้างาน = บรรทัดบน · คำที่ไอทีใช้ = บรรทัดล่าง")
node(s, 0.5, 1.9, 2.3, 1.3, "ผู้ใช้ ~70 คน", ["PC · มือถือ · จอ TV"])
node(s, 3.45, 1.9, 2.8, 1.3, "หน้าเว็บ", ["Static site (React + Vite)", "อยู่บน Render.com"])
T.rect(s, 6.85, 1.65, 6.0, 3.75, fill=GREEN_TINT, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T.text(s, 7.0, 1.7, 5.7, 0.4, [[("Supabase (คลาวด์) — 2 project", {"sz": 14, "b": True, "c": GREEN})]])
node(s, 7.05, 2.15, 2.75, 1.25, "ฐานข้อมูล Main", ["PostgreSQL · 277 MB", "คน · สิทธิ์ · 4M"])
node(s, 9.95, 2.15, 2.75, 1.25, "ฐานข้อมูล DR", ["PostgreSQL · 198 MB", "ยอดผลิต · เครื่องหยุด"])
for i, (h, d) in enumerate([("ล็อกอิน", "Auth (JWT)"), ("จอขยับเอง", "Realtime (WebSocket)"),
                            ("หลังบ้าน 24 ตัว", "Edge Functions (Deno)"), ("ไฟล์/รูป 212 MB", "Object storage (S3)")]):
    x = 7.05 + (i % 2) * 2.9; y = 3.55 + (i // 2) * 0.9
    T.rect(s, x, y, 2.75, 0.78, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, x, y, 2.75, 0.78, [[(h, {"sz": 13, "b": True, "c": GREEN_D})], [(d, {"sz": 11, "c": GREY})]],
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sa=0)
arrow(s, 2.8, 2.55, 3.45, 2.55); arrow(s, 6.25, 2.55, 7.05, 2.75)
T.text(s, 0.5, 3.55, 6.1, 0.4, [[("ระบบต้องคุยกับภายนอก", {"sz": 15, "b": True, "c": ORANGE})]])
for i, (h, d) in enumerate([("Telegram", "ส่งแจ้งเตือน + รับข้อความตอบ (webhook เข้า)"),
                            ("Push มือถือ", "ผ่าน push service ของ Google / Apple"),
                            ("ไฟล์ EDI ลูกค้า", "สคริปต์ดึงจาก Outlook ส่งเข้าระบบ"),
                            ("งานตามเวลา", "pg_cron 12 งาน + pg_net ยิงหาหลังบ้าน")]):
    y = 4.0 + i * 0.5
    T.rect(s, 0.5, y + 0.12, 0.12, 0.26, fill=ORANGE)
    T.text(s, 0.75, y, 5.9, 0.5, [[(h + "  ", {"sz": 13, "b": True, "c": GREEN_D}), (d, {"sz": 13, "c": GREY})]], anchor=MSO_ANCHOR.MIDDLE)
T.rect(s, 0.5, 6.15, 12.35, 0.6, fill=GREEN_D, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T.text(s, 0.7, 6.15, 12.0, 0.6, [[("ข้อมูลรวมไม่ถึง 1 GB · ผู้ใช้ราว 70 คน  ·  ", {"sz": 15, "b": True, "c": WHITE}),
                                 ("ระบบเล็ก ไม่ต้องใช้เครื่องแรง", {"sz": 15, "b": True, "c": WHITE})]], anchor=MSO_ANCHOR.MIDDLE)
notes(s, "ชี้ทีละกล่อง: คนเปิดหน้าเว็บ หน้าเว็บคุยกับ Supabase ซึ่งมีฐานข้อมูล 2 ก้อน + ล็อกอิน + realtime + หลังบ้าน 24 ตัว + ที่เก็บไฟล์ ด้านซ้ายล่างคือสิ่งที่ระบบต้องออกไปคุยข้างนอก ตัวเลขขนาดวัดไว้ 12 ต.ค. โค้ดทั้งหมดอยู่บน GitHub มีประวัติทุกการแก้ (migration 622 ไฟล์)")

# ── 3 ทางเลือก ──
s = content("ทางเลือกการย้าย", "ขอให้ไอทีช่วยตัดสิน — ฝั่งผู้ใช้เสนอ A")
opts = [("A", "ติดตั้ง Supabase บนเซิร์ฟเวอร์บริษัท", "self-host ด้วย Docker + nginx เสิร์ฟหน้าเว็บ",
         "แทบไม่แก้ — เปลี่ยน URL / key", ["ใกล้ของเดิมที่สุด", "ข้อมูลอยู่ในบ้าน"], ["ไอทีดูแล Docker เอง", "backup · อัปเดต · SSL"], True),
        ("B", "คงคลาวด์ ให้ไอทีคุมบัญชี", "โอนบัญชีเป็นของบริษัท · จ่ายแพ็กเกจ Pro",
         "ไม่ต้องแก้", ["เร็วที่สุด", "Supabase ทำ backup ให้"], ["ข้อมูลยังอยู่นอกบริษัท", "มีค่ารายเดือน"], False),
        ("C", "เขียนใหม่ลงฐานข้อมูลที่บริษัทใช้", "เช่น SQL Server",
         "แก้หนักมาก", ["ใช้ระบบเดียวกับที่ไอทีคุ้นเคย"], ["ต้องทำล็อกอิน realtime หลังบ้านใหม่", "ใช้เวลาหลายเดือน"], False)]
for i, (k, h, how, code, pro, con, rec) in enumerate(opts):
    x = 0.5 + i * 4.15; y = 1.75; w = 4.0; hh = 4.55
    T.rect(s, x, y, w, hh, fill=WHITE, line=(ORANGE if rec else BORDER), lw=(2.5 if rec else 0.75),
           shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    T.text(s, x + 0.25, y + 0.15, 0.8, 0.8, [[(k, {"sz": 40, "b": True, "c": ORANGE if rec else GREEN})]])
    if rec:
        T.rect(s, x + 2.55, y + 0.25, 1.25, 0.4, fill=ORANGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        T.text(s, x + 2.55, y + 0.25, 1.25, 0.4, [[("ข้อเสนอ", {"sz": 13, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    T.text(s, x + 0.25, y + 0.9, w - 0.5, 1.1, [[(h, {"sz": 16, "b": True, "c": GREEN_D})], [(how, {"sz": 12, "c": GREY})]], sa=3)
    T.text(s, x + 0.25, y + 2.05, w - 0.5, 0.4, [[("แก้โค้ด: ", {"sz": 13, "b": True, "c": GREEN_D}), (code, {"sz": 13, "c": GREEN_D})]])
    T.text(s, x + 0.25, y + 2.5, w - 0.5, 0.95, [[("ข้อดี: " + t, {"sz": 13, "c": GREEN})] for t in pro], sa=3)
    T.text(s, x + 0.25, y + 3.45, w - 0.5, 1.0, [[("ข้อเสีย: " + t, {"sz": 13, "c": ORANGE})] for t in con], sa=3)
T.text(s, 0.5, 6.4, 12.3, 0.45, [[("ถ้าไอทียังไม่พร้อมดูแล Docker: ใช้ B ชั่วคราวไปก่อน แล้วค่อยย้ายเป็น A", {"sz": 14, "b": True, "c": GREEN_D})]])
notes(s, "เสนอ A เพราะโค้ดแทบไม่ต้องแก้ แค่เปลี่ยนที่อยู่ของเซิร์ฟเวอร์ แต่ภาระอยู่ที่ไอที: ดูแล Docker ทำ backup อัปเดต ทำใบรับรอง SSL  B คือถ้าไอทียังไม่พร้อม ให้บริษัทถือบัญชีคลาวด์ไว้ก่อน  C ไม่แนะนำ เพราะต้องเขียนส่วนล็อกอิน realtime หลังบ้านใหม่ทั้งหมด")

# ── 4 ต้องเตรียม ──
s = content("ถ้าย้ายเข้าบ้าน (ทางเลือก A) ต้องเตรียมอะไร", "ตัวเลขสเปกเป็นค่าเริ่มต้น — ไอทียืนยันจากเอกสาร Supabase อีกครั้ง")
cards = [("เครื่อง", ["~4 vCPU · 8 GB RAM", "ดิสก์ 100 GB (เผื่อโต)", "รัน Docker ได้"]),
         ("เน็ตออก", ["Telegram API", "push service Google / Apple", "GitHub (ดึงโค้ดไป build)"]),
         ("เน็ตเข้า", ["Telegram webhook ต้องมี HTTPS", "ใช้จากมือถือนอกโรงงานไหม?", "ต้องมีโดเมน + SSL"]),
         ("ดูแลต่อเนื่อง", ["pg_dump รายวัน + backup ไฟล์", "อัปเดต Supabase เป็นระยะ", "deploy อัตโนมัติเมื่อโค้ดใหม่"])]
for i, (h, ls) in enumerate(cards):
    x = 0.5 + i * 3.1
    T.rect(s, x, 1.75, 2.95, 2.75, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    T.rect(s, x, 1.75, 2.95, 0.5, fill=GREEN_D)
    T.text(s, x, 1.75, 2.95, 0.5, [[(h, {"sz": 16, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    T.text(s, x + 0.2, 2.4, 2.6, 2.0, [[("•  " + l, {"sz": 13, "c": GREEN})] for l in ls], sa=8)
T.rect(s, 0.5, 4.75, 12.35, 1.65, fill=WHITE, line=ORANGE, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T.rect(s, 0.5, 4.75, 0.14, 1.65, fill=ORANGE)
T.text(s, 0.85, 4.85, 11.8, 1.5, [
    [("ความปลอดภัย — บอกตรงๆ", {"sz": 16, "b": True, "c": ORANGE})],
    [("ฐานข้อมูล DR ตอนนี้หน้าเว็บเข้าด้วยกุญแจสาธารณะ (anon key) ยังไม่ผูกกับการล็อกอิน — เป็นจุดอ่อนที่รู้อยู่และมีแผนแก้", {"sz": 13, "c": GREEN_D})],
    [("ถ้าย้ายเข้าวงในบริษัทและไม่เปิดออกอินเทอร์เน็ต ความเสี่ยงข้อนี้ลดลงมาก", {"sz": 13, "c": GREEN_D})]], sa=4)
notes(s, "ไล่ 4 กล่อง: เครื่อง, ต้องออกเน็ตไปไหน, ต้องมีอะไรเข้ามาจากข้างนอก, งานดูแลต่อเนื่อง กล่องส้มคือเรื่องความปลอดภัย พูดตรงๆ ว่ามีจุดอ่อนที่รู้อยู่ ไอทีจะถามเรื่องนี้แน่นอน")

# ── 5 ภาพเมื่อขยายไปหลายโรงงาน ──
s = content("ภาพเมื่อขยายไปหลายโรงงานในกลุ่ม", "ศูนย์กลางเดียวที่ INET · ทุกโรงงานใช้โปรแกรมชุดเดียวกัน · ข้อมูลแยกกันด้วยรหัสบริษัท")
# ฝั่งโรงงาน
T.text(s, 0.4, 1.5, 3.6, 0.35, [[("โรงงาน (เพิ่มได้เรื่อยๆ)", {"sz": 14, "b": True, "c": GREEN})]])
plants = [("TSAT4 (เรา)", "ใช้งานจริง", True), ("โรงงานในไทย", "TSAT1 · TSLA …", False), ("ต่างประเทศ", "เวียดนาม · อินเดีย …", False)]
for i, (h, d, me) in enumerate(plants):
    y = 1.9 + i * 1.45
    T.rect(s, 0.4, y, 3.6, 1.3, fill=WHITE, line=(ORANGE if me else BORDER), lw=(2 if me else 0.75), shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    T.text(s, 0.55, y + 0.05, 3.3, 0.45, [[(h, {"sz": 14, "b": True, "c": ORANGE if me else GREEN_D}), ("  " + d, {"sz": 11, "c": GREY})]], anchor=MSO_ANCHOR.MIDDLE)
    for j, t in enumerate(["PC · จอ TV", "มือถือ", "กล่อง Edge"]):
        x = 0.55 + j * 1.15
        T.rect(s, x, y + 0.55, 1.05, 0.6, fill=GREEN_TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        T.text(s, x, y + 0.55, 1.05, 0.6, [[(t, {"sz": 11, "b": True, "c": GREEN_D})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ชั้นเครือข่าย
for i, (h, d, y) in enumerate([("VPN ระหว่างไซต์ / IP whitelist", "PC · จอ TV · กล่อง Edge", 2.2),
                               ("Cloudflare Zero Trust + SSO", "มือถือ · ล็อกอินด้วยบัญชี M365", 3.95)]):
    T.rect(s, 4.45, y, 2.55, 1.35, fill=WHITE, line=GREEN, lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, 4.55, y + 0.05, 2.35, 1.25, [[(h, {"sz": 13, "b": True, "c": GREEN_D})], [(d, {"sz": 11, "c": GREY})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sa=3)
arrow(s, 4.0, 2.55, 4.45, 2.8); arrow(s, 4.0, 4.0, 4.45, 3.0); arrow(s, 4.0, 5.45, 4.45, 3.2)
arrow(s, 4.0, 2.7, 4.45, 4.5); arrow(s, 4.0, 4.15, 4.45, 4.6); arrow(s, 4.0, 5.6, 4.45, 4.7)
arrow(s, 7.0, 2.85, 7.4, 2.85); arrow(s, 7.0, 4.6, 7.4, 4.6)
# ศูนย์กลาง INET
T.rect(s, 7.4, 1.5, 5.5, 4.65, fill=GREEN_TINT, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T.text(s, 7.55, 1.55, 5.2, 0.4, [[("ศูนย์กลาง ESM ที่ INET — Kubernetes 3 เครื่อง (ล่ม 1 ยังวิ่ง)", {"sz": 13, "b": True, "c": GREEN})]])
cen = [("หน้าเว็บ", "image เดียว ทุกโรงงาน"), ("ล็อกอิน + สิทธิ์", "SSO · role ต่อบริษัท"),
       ("ฐานข้อมูล Main", "คน · ผังองค์กร TSG→บริษัท"), ("ฐานข้อมูล DR", "ข้อมูลผลิต · แยกด้วย company_id"),
       ("หลังบ้าน + cron", "แจ้งเตือน · สแกน PM"), ("รายงานรวมกลุ่ม", "OBEYA · Group Overview")]
for k, (h, d) in enumerate(cen):
    x = 7.55 + (k % 2) * 2.65; y = 2.0 + (k // 2) * 0.98
    T.rect(s, x, y, 2.55, 0.85, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.text(s, x, y, 2.55, 0.85, [[(h, {"sz": 13, "b": True, "c": GREEN_D})], [(d, {"sz": 11, "c": GREY})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sa=0)
T.rect(s, 7.55, 4.98, 5.2, 1.05, fill=WHITE, line=ORANGE, lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T.text(s, 7.7, 4.98, 4.95, 1.05, [[("Backup นอกไซต์ (data center ที่ 2)", {"sz": 13, "b": True, "c": ORANGE})],
                                  [("ฐานข้อมูล + ไฟล์ทุกวัน · ทดลองกู้คืนจริงปีละครั้ง", {"sz": 11, "c": GREEN_D})]], anchor=MSO_ANCHOR.MIDDLE, sa=2)
T.text(s, 4.3, 5.4, 2.95, 0.9, [[("กล่อง Edge: ", {"sz": 11, "b": True, "c": GREEN_D}), ("รับสัญญาณเครื่องจักร + ไฟล์ EDI ส่งเข้าศูนย์", {"sz": 11, "c": GREY})],
                                [("ภายนอก: ", {"sz": 11, "b": True, "c": GREEN_D}), ("Telegram · Push · GitHub", {"sz": 11, "c": GREY})]], sa=3)
T.rect(s, 0.4, 6.35, 12.5, 0.45, fill=GREEN_D, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T.text(s, 0.55, 6.35, 12.2, 0.45, [[("เพิ่มโรงงาน = เพิ่มรหัสบริษัท + ผู้ใช้ + ทะเบียน  ·  ไม่ต้องตั้งเซิร์ฟเวอร์ใหม่ทุกโรงงาน", {"sz": 14, "b": True, "c": WHITE})]], anchor=MSO_ANCHOR.MIDDLE)
notes(s, "ภาพปลายทาง: ทุกโรงงานในกลุ่มต่อเข้าศูนย์กลางเดียวที่ INET ใช้โปรแกรมชุดเดียวกัน ข้อมูลแต่ละบริษัทแยกด้วยรหัสบริษัท (company_id) ผู้บริหารกลุ่มดูรวมได้จากจอ OBEYA/Group Overview  PC จอ TV และกล่อง Edge ต่อผ่าน VPN ระหว่างไซต์ มือถือเข้าผ่าน Cloudflare + ล็อกอินบัญชี M365 ตัวเดียวกับอีเมล ศูนย์กลางควรเป็น 3 เครื่อง ไม่ใช่เครื่องเดียว เพราะถ้าล่มจะหยุดทุกโรงงานพร้อมกัน และต้องมี backup นอกไซต์  ข้อยกเว้น: ถ้าประเทศไหนกฎหมายบังคับให้ข้อมูลอยู่ในประเทศ ค่อยแยกชุดให้โรงงานนั้น")

# ── 6 ขยายทีละขั้น ──
s = content("ขยายทีละขั้น · อะไรใช้ร่วม อะไรแยกต่อโรงงาน")
phases = [("ขั้น 1", "TSAT4 ย้ายขึ้น INET", ["เครื่องตามผังไอที", "pipeline อัปเดตอัตโนมัติ", "backup + มือถือเข้าได้"], True),
          ("ขั้น 2", "เตรียมรองรับหลายบริษัท", ["ใส่ company_id + สิทธิ์ต่อบริษัท", "ปิดช่องโหว่ฐานข้อมูล DR", "SSO ด้วยบัญชี M365"], False),
          ("ขั้น 3", "โรงงานนำร่องที่ 2", ["VPN + กล่อง Edge ที่ไซต์", "ขยายเป็น 3 เครื่อง", "แล้วค่อยเพิ่มทีละโรงงาน"], False)]
for i, (k, h, ls, now) in enumerate(phases):
    x = 0.5 + i * 4.15
    T.rect(s, x, 1.55, 4.0, 2.55, fill=WHITE, line=(ORANGE if now else BORDER), lw=(2 if now else 0.75), shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
    T.rect(s, x, 1.55, 4.0, 0.5, fill=(ORANGE if now else GREEN_D))
    T.text(s, x, 1.55, 4.0, 0.5, [[(k + "  ·  " + h, {"sz": 15, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    T.text(s, x + 0.25, 2.2, 3.5, 1.8, [[("•  " + l, {"sz": 13, "c": GREEN})] for l in ls], sa=8)
for j, (h, ls, c) in enumerate([("ใช้ร่วมทุกโรงงาน", ["โปรแกรม + สูตรคำนวณ (OEE ฯลฯ) ชุดเดียว", "เซิร์ฟเวอร์ · backup · ทีมดูแล", "ล็อกอิน SSO · จอรวมกลุ่ม"], GREEN),
                                ("แยกต่อโรงงาน", ["ข้อมูลผลิต / คน (ด้วย company_id)", "ปฏิทิน · เวลาพัก · เป้า KPI · ผังโรงงาน", "เลขฟอร์มเอกสาร · ห้อง Telegram"], ORANGE)]):
    x = 0.5 + j * 6.25
    T.rect(s, x, 4.35, 6.1, 1.85, fill=GREEN_TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T.rect(s, x, 4.35, 0.14, 1.85, fill=c)
    T.text(s, x + 0.35, 4.42, 5.6, 1.75, [[(h, {"sz": 15, "b": True, "c": c})]] + [[("•  " + l, {"sz": 13, "c": GREEN_D})] for l in ls], sa=5)
T.text(s, 0.5, 6.35, 12.3, 0.45, [[("ต้องทำขั้น 2 ก่อนรับโรงงานที่ 2 — ไม่งั้นโรงงานหนึ่งจะมองเห็นข้อมูลของอีกโรงงานได้", {"sz": 14, "b": True, "c": ORANGE})]])
notes(s, "ไม่ต้องรอทำทุกอย่างพร้อมกัน ขั้น 1 คือย้าย TSAT4 ตามผังที่ไอทีเสนอ ขั้น 2 คืองานฝั่งโปรแกรมให้รองรับหลายบริษัท ข้อนี้สำคัญ: ตอนนี้ฐานข้อมูล DR ไม่ผูกกับการล็อกอิน ถ้าเอาหลายบริษัทมาอยู่ฐานเดียวโดยไม่แก้ บริษัทหนึ่งจะเห็นข้อมูลอีกบริษัท ขั้น 3 ค่อยรับโรงงานนำร่อง ส่วนล่างบอกว่าอะไรใช้ร่วม อะไรแยก  ระบบออกแบบไว้รองรับแล้ว (ผังองค์กรเพิ่มชั้นบริษัทได้ ตารางรายวันไม่ต้องแก้ เพราะรู้บริษัทจากไลน์)")

# ── 7 คำถาม + ขั้นต่อไป ──
s = content("ขอคำตอบจากไอที + ขั้นต่อไป")
qs = ["มีเครื่องที่รัน Docker ได้ไหม / ใช้ VM ตัวไหน",
      "ใช้เฉพาะในโรงงาน หรือต้องเข้าจากมือถือนอกบริษัทด้วย",
      "มีโดเมนบริษัทให้ใช้ไหม · ใครออกใบรับรอง SSL",
      "นโยบาย backup / เก็บข้อมูลของบริษัท",
      "ใครเป็นผู้ดูแลเครื่อง · ใครได้สิทธิ์ admin",
      "เครื่องในบริษัทออกไป Telegram / push มือถือได้ไหม"]
T.rect(s, 0.5, 1.6, 7.3, 4.85, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
T.text(s, 0.75, 1.7, 6.9, 0.45, [[("คำถามถึงไอที", {"sz": 17, "b": True, "c": GREEN_D})]])
for i, q in enumerate(qs):
    y = 2.25 + i * 0.68
    T.rect(s, 0.75, y + 0.06, 0.42, 0.42, fill=GREEN, shape=MSO_SHAPE.OVAL)
    T.text(s, 0.75, y + 0.06, 0.42, 0.42, [[(str(i + 1), {"sz": 14, "b": True, "c": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    T.text(s, 1.35, y, 6.3, 0.55, [[(q, {"sz": 14, "c": GREEN_D})]], anchor=MSO_ANCHOR.MIDDLE)
T.rect(s, 8.05, 1.6, 4.8, 2.6, fill=GREEN_TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T.text(s, 8.3, 1.7, 4.4, 2.45, [[("ขั้นต่อไป", {"sz": 17, "b": True, "c": GREEN_D})]] + [
    [("•  " + t, {"sz": 13, "c": GREEN})] for t in
    ["ยืนยันผังไอที + ตอบคำถามข้างซ้าย", "ทำแผนย้ายทีละขั้น + วิธีย้อนกลับ", "ทดลองคู่ขนานก่อนสลับจริง",
     "ย้ายจริง หลัง roadshow ปลาย ต.ค."]], sa=6)
T.rect(s, 8.05, 4.4, 4.8, 2.05, fill=WHITE, line=ORANGE, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T.text(s, 8.3, 4.5, 4.4, 1.9, [[("ยังไม่รับปาก", {"sz": 17, "b": True, "c": ORANGE})]] + [
    [("•  " + t, {"sz": 13, "c": GREEN_D})] for t in
    ["วันย้ายเสร็จ — รอไอทีประเมิน", "ย้ายแล้ว 'ไม่ต้องแก้อะไรเลย'", "ค่าใช้จ่าย — ขึ้นกับเครื่องที่มี"]], sa=6)
notes(s, "อ่านคำถามทีละข้อ จดคำตอบ ข้อ 2 สำคัญที่สุด เพราะกำหนดเรื่อง firewall และ SSL ทั้งหมด ปิดด้วยขั้นต่อไป และย้ำว่าไม่ย้ายก่อน roadshow")

p.save(OUT)
print("saved", OUT, len(p.slides), "slides")
